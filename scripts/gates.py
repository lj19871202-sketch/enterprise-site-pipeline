#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""企业官网资料包 · 质量门禁

一条命令跑完所有"机器可判定"的检查，任一 error 级不通过即非零退出。

    python scripts/gates.py --config pipeline.json
    python scripts/gates.py --config pipeline.json --only en_pinyin,docx_sync
    python scripts/gates.py --config pipeline.json --json build/gates.json

门禁是断言，不是报告。红了先修数据，不要下调阈值迁就数据。
阈值与门禁清单见 references/quality-gates.md。
"""
import argparse
import collections
import glob
import hashlib
import json
import os
import re
import sys

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

DIRS2 = {
    "factory": "1.企业工厂图",
    "product": "2.企业产品图",
    "logo": "3.企业logo",
    "cert": "4.资质证书",
}
EN_KEYS = ["英文名", "英文标题", "品牌", "英文简介", "产品英名", "子品类"]
CJK = re.compile(r"[\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff（）【】「」，。！？、；：]")
LATIN = re.compile(r"[A-Za-z]")
# 至少 3 个连续英文词，避免型号（WMS/GaN/5A）被误当成英文译文。
EN_SENTENCE = re.compile(r"[A-Za-z]{2,}(?:\s+[A-Za-z][A-Za-z'’,.-]*){2,}")
NOISE = re.compile(
    r"首页[»>]|您当前的位置|欢迎光临|信息纠错|客服中心|会员级别|顺企|友情链接|"
    r"荟萃网库|未经核实|询盘|贸易通|立即注册|请发送您要|营业执照号码|经营范围|"
    r"技术支持|热门关键词|机电局|机电之家|瀚睿|Powered by|Toggle navigation|"
    r"Jump to main|Sign in|Read ?more|Copyright|ICP备|备案号"
)
PLACEHOLDER = re.compile(r"网站建设中|敬请期待|Lorem ipsum", re.I)

# 常见英文词（门禁只用于"可疑度"启发式，不追求完整）
COMMON = set("""
the of and to in for with a an on by from is are was were be been being as at it its
this that these those their our your his her company co ltd limited group main product
products service services include includes including mainly engaged located province
city area covers square kilometers more details please visit official website also
offer offers provide provides supply wholesale retail quality inspection process
production sale sales manufacture manufacturing technology new high clean energy
equipment system systems parts component components machinery export import customers
clients worldwide domestic international research development design development
""".split())
SUFFIXES = (
    "tion", "sion", "ment", "ness", "ance", "ence", "able", "ible", "ings", "ing",
    "ies", "ied", "ers", "ors", "ed", "es", "ly", "ial", "ical", "ic", "ous",
    "ive", "ity", "ate", "ize", "ise", "ent", "ant", "ary", "ory", "al",
)


def expand(p):
    return os.path.expanduser(p) if isinstance(p, str) else p


def host_of(url):
    from urllib.parse import urlparse
    try:
        return urlparse(str(url or "")).netloc.lower().split("@")[-1].split(":")[0]
    except Exception:
        return ""


def same_host(a, b):
    return bool(a and b and (a == b or a.endswith("." + b) or b.endswith("." + a)))


def norm_text(s):
    return re.sub(r"\s+", " ", str(s or "")).strip()


def company_dir_name(name):
    """与 local_pipeline.safe_filename 保持一致。"""
    s = norm_text(name)
    s = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", s)
    s = s.strip(" .")
    return s[:120] or "company"


def company_path(root, name):
    exact = os.path.join(root, name)
    if os.path.isdir(exact):
        return exact
    safe = os.path.join(root, company_dir_name(name))
    return safe if os.path.isdir(safe) else exact


def read_json(path):
    with open(path, encoding="utf-8") as fh:
        return json.load(fh)


def load_company_files(raw_dir):
    out = {}
    for fp in glob.glob(os.path.join(raw_dir, "*.json")):
        try:
            d = read_json(fp)
        except Exception as exc:
            print(f"  ! 跳过无法解析的档案 {fp}: {exc}", file=sys.stderr)
            continue
        if d.get("名称"):
            out[d["名称"]] = d
    return out


def docx_paragraphs(path):
    from docx import Document
    return [p.text.strip() for p in Document(path).paragraphs if p.text.strip()]


def split_docx(ps):
    """返回 (标题, 中文段列表, 英文段列表)：按是否含中日韩字符归类。"""
    return ps[0], [t for t in ps[1:] if CJK.search(t)], [t for t in ps[1:] if not CJK.search(t)]


def suspect_ratio(paragraph):
    """段内疑似拼音/栏目词的 token 占比。"""
    toks = [t for t in re.findall(r"[A-Za-z]+", paragraph) if len(t) >= 4]
    if not toks:
        return 0.0
    bad = 0
    for t in toks:
        w = t.lower()
        if len(w) <= 7 or w in COMMON or any(w.endswith(s) for s in SUFFIXES):
            continue
        bad += 1
    return bad / len(toks)


class Gate:
    def __init__(self, gid, level, ok, detail, offenders=None):
        self.id, self.level, self.ok = gid, level, ok
        self.detail, self.offenders = detail, offenders or []

    def as_dict(self):
        return {
            "id": self.id, "level": self.level, "ok": self.ok,
            "detail": self.detail, "offenders": self.offenders[:20],
            "offender_count": len(self.offenders),
        }


# ---------------------------------------------------------------- 各门禁实现

def gate_site_discovery(ctx):
    """官网必须找到；自动发现低/中置信度必须由 Codex 打开确认。"""
    bad, warn = [], []
    for name, d in sorted(ctx["companies"].items()):
        site = str(d.get("官网") or "").strip()
        conf = str(d.get("置信度") or "").strip()
        if not site:
            bad.append(f"{name}: 未发现官网")
        elif conf.startswith("低") or "需复核" in conf:
            warn.append(f"{name}: 官网自动发现置信度低，Codex 需打开确认 {site}")
        elif conf.startswith("中"):
            warn.append(f"{name}: 官网自动发现置信度中 {site}")
    if bad:
        return Gate("site_discovery", "error", False,
                    f"官网未确认 {len(bad)} 家；需复核 {len(warn)} 家", bad + warn)
    if warn:
        return Gate("site_discovery", "warn", True,
                    f"官网均已找到；{len(warn)} 家置信度中/低，Codex 需打开确认", warn)
    return Gate("site_discovery", "error", True,
                f"{len(ctx['companies'])} 家官网已确认", [])


def gate_coverage(ctx):
    raw, root, en, summary = ctx["raw"], ctx["deliverable"], ctx["en"], ctx["summary"]
    raws = ctx["companies"]
    dirs = set(d for d in os.listdir(root)
               if os.path.isdir(os.path.join(root, d))) if os.path.isdir(root) else set()
    enk = set(en) if isinstance(en, dict) else set()
    rows = ctx["summary_rows"]
    bad = []
    if set(raws) - dirs:
        bad.append(f"有档案无目录 {sorted(set(raws) - dirs)[:5]}")
    if dirs - set(raws):
        bad.append(f"有目录无档案 {sorted(dirs - set(raws))[:5]}")
    if set(raws) - enk:
        bad.append(f"有档案无英文 {sorted(set(raws) - enk)[:5]}")
    if enk - set(raws):
        bad.append(f"有英文无档案 {sorted(enk - set(raws))[:5]}")
    summary_names = ctx.get("summary_names")
    if summary_names is not None:
        sm = set(summary_names)
        if set(raws) - sm:
            bad.append(f"有档案无汇总 {sorted(set(raws) - sm)[:5]}")
        if sm - set(raws):
            bad.append(f"有汇总无档案 {sorted(sm - set(raws))[:5]}")
    exp = ctx["expected"]
    if exp and len(raws) != exp:
        bad.append(f"企业数 {len(raws)} != 期望 {exp}")
    if rows is not None and rows != len(raws):
        bad.append(f"汇总表 {rows} 行 != 档案 {len(raws)} 家")
    return Gate("coverage", "error", not bad, "; ".join(bad) or f"{len(raws)} 家四方一致",
                sorted(dirs - set(raws)))


def gate_structure(ctx):
    bad = []
    for name in sorted(ctx["companies"]):
        home = os.path.join(ctx["deliverable"], name)
        if not os.path.isdir(home):
            continue
        missing = [d for d in DIRS2.values() if not os.path.isdir(os.path.join(home, d))]
        if not glob.glob(os.path.join(home, "5.企业介绍", "*.docx")):
            missing.append("简介docx")
        if not os.path.isfile(os.path.join(home, "产品清单.xlsx")):
            missing.append("产品清单.xlsx")
        if missing:
            bad.append(f"{name}:{','.join(missing)}")
    return Gate("structure", "error", not bad, f"{len(bad)} 家结构缺失", bad)


def gate_images(ctx):
    dead, empty, md5, orphan = [], [], collections.defaultdict(list), []
    for name, d in ctx["companies"].items():
        home = company_path(ctx["deliverable"], name)
        if not any(d.get(cat) for cat in DIRS2):
            empty.append(name)
        for cat, folder in DIRS2.items():
            recorded = set()
            for it in d.get(cat) or []:
                fn = it.get("file", "")
                recorded.add(fn)
                fp = os.path.join(home, folder, fn)
                if not (fn and os.path.isfile(fp)):
                    dead.append(f"{name}/{cat}/{fn}")
                    continue
                with open(fp, "rb") as fh:
                    md5[hashlib.md5(fh.read()).hexdigest()].append(f"{name}/{folder}")
            folder_path = os.path.join(home, folder)
            if os.path.isdir(folder_path):
                for fn in sorted(os.listdir(folder_path)):
                    if os.path.isfile(os.path.join(folder_path, fn)) and fn not in recorded:
                        orphan.append(f"{name}/{folder}/{fn}")
    cross = [v for v in md5.values() if len({x.split("/")[1] for x in v}) > 1]
    ok = not dead and not cross and not empty and not orphan
    detail = (f"失效引用 {len(dead)}、跨类重复组 {len(cross)}、四类全空 {len(empty)} 家、"
              f"未被 raw 记录的孤儿图 {len(orphan)} 条")
    offenders = dead + [str(c) for c in cross[:5]] + [f"{n}:四类图片全空" for n in empty] + orphan
    return Gate("images", "error", ok, detail, offenders)


def gate_image_required(ctx):
    """logo 与工厂图是交付重点，空类必须显式告警，但不硬性阻断无实拍的企业。"""
    missing = []
    for name, d in sorted(ctx["companies"].items()):
        for cat, label in (("logo", "企业logo"), ("factory", "企业工厂图")):
            if not d.get(cat):
                missing.append(f"{name}:{label}为空")
    return Gate(
        "image_required", "warn", not missing,
        f"logo/工厂图为空 {len(missing)} 条（需人工补图或在交付说明中标记数据边界）",
        missing,
    )


def gate_image_provenance(ctx):
    bad, offsite = [], []
    for name, d in sorted(ctx["companies"].items()):
        site_host = host_of(d.get("官网", ""))
        for cat, folder in DIRS2.items():
            for it in d.get(cat) or []:
                label = f"{name}/{folder}/{it.get('file', '?')}"
                frm = it.get("from", "")
                fh = host_of(frm)
                if not frm:
                    bad.append(f"{label}: 缺来源页面")
                elif not same_host(fh, site_host):
                    bad.append(f"{label}: 来源页域名 {fh or frm[:40]} 与官网不一致")
                uh = host_of(it.get("url", ""))
                if uh and site_host and not same_host(uh, site_host):
                    offsite.append(f"{label}: 图片直链 {uh}")
    detail = f"来源页非官网 {len(bad)} 条；图片直链外域 {len(offsite)} 条（CDN 需人工确认）"
    return Gate("image_provenance", "error", not bad, detail, bad + offsite)


def gate_en_entry(ctx):
    bad = []
    for name, v in (ctx["en"] or {}).items():
        probs = [f"缺{k}" for k in EN_KEYS if k not in v]
        for k in ("英文名", "英文标题", "品牌"):
            if k not in v:
                continue
            if not isinstance(v.get(k), str) or not v[k].strip():
                probs.append(f"{k}为空")
        ps = v.get("英文简介")
        if not isinstance(ps, list) or len(ps) != 3:
            probs.append(f"简介{len(ps) if isinstance(ps, list) else 0}段")
            ps = ps if isinstance(ps, list) else []
        for p in ps:
            if not isinstance(p, str) or len(p.strip()) < 60:
                probs.append("段落过短")
        pm = v.get("产品英名")
        if "产品英名" in v and not isinstance(pm, dict):
            probs.append("产品英名非映射")
        if "子品类" in v and not isinstance(v.get("子品类"), str):
            probs.append("子品类非文本")
        texts = []
        for f in ("英文名", "英文标题", "品牌"):
            if isinstance(v.get(f), str):
                texts.append(v[f])
        texts += [p for p in ps if isinstance(p, str)]
        if isinstance(pm, dict):
            texts += [x for x in pm.values() if isinstance(x, str)]
        if any(CJK.search(s) for s in texts):
            probs.append("含中文/全角")
        if probs:
            bad.append(f"{name}:{'/'.join(sorted(set(probs)))}")
    return Gate("en_entry", "error", not bad, f"{len(bad)} 家英文条目不完整", bad)


def gate_en_ascii(ctx):
    bad, worst = [], 0.0
    for name, v in (ctx["en"] or {}).items():
        for i, p in enumerate(v.get("英文简介") or []):
            if not p:
                continue
            ratio = sum(1 for c in p if ord(c) > 127) / len(p)
            worst = max(worst, ratio)
            if ratio > ctx["ascii_max"]:
                bad.append(f"{name}#{i + 1} 非ASCII {ratio:.1%}")
    return Gate("en_ascii", "error", not bad,
                f"{len(bad)} 段非ASCII超阈值；最高 {worst:.1%}", bad)


def gate_en_pinyin(ctx):
    bad = []
    for name, v in (ctx["en"] or {}).items():
        for i, p in enumerate(v.get("英文简介") or []):
            r = suspect_ratio(p)
            if r >= ctx["pinyin_max"]:
                bad.append(f"{name}#{i + 1} 可疑 {r:.0%}")
    return Gate("en_pinyin", "warn", not bad,
                f"{len(bad)} 段疑似拼音/栏目词（启发式，需人工确认）", bad)


def gate_product_en(ctx):
    import openpyxl
    bad = []
    for fp in glob.glob(os.path.join(ctx["deliverable"], "*", "产品清单.xlsx")):
        name = os.path.basename(os.path.dirname(fp))
        try:
            ws = openpyxl.load_workbook(fp, read_only=True)["产品清单"]
        except Exception as exc:
            bad.append(f"{name}:无法读取 {exc}")
            continue
        hdr = [c.value for c in next(ws.iter_rows(max_row=1))]
        col = hdr.index("产品名称(中英文)") if "产品名称(中英文)" in hdr else 2
        for r, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            cell = row[col] if col < len(row) else None
            if cell is None or not str(cell).strip():
                continue
            text = str(cell).strip()
            _zh, sep, eng = text.partition(" / ")
            if not sep or not re.search(r"[A-Za-z]{2,}", eng):
                bad.append(f"{name} 行{r}: {text[:30]}")
    return Gate("product_en", "error", not bad, f"{len(bad)} 行产品缺英文", bad)


def gate_product_detail(ctx):
    """产品详情列必须中英双语；官网真实详情覆盖率作为告警指标。"""
    import openpyxl
    bad, empty, total, hit = [], 0, 0, 0
    for fp in glob.glob(os.path.join(ctx["deliverable"], "*", "产品清单.xlsx")):
        name = os.path.basename(os.path.dirname(fp))
        try:
            ws = openpyxl.load_workbook(fp, read_only=True)["产品清单"]
        except Exception as exc:
            bad.append(f"{name}:无法读取 {exc}")
            continue
        hdr = [c.value for c in next(ws.iter_rows(max_row=1))]
        col = hdr.index("产品详情") if "产品详情" in hdr else 6
        for r, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            text = norm_text(row[col] if col < len(row) else None)
            total += 1
            if not text:
                empty += 1
                continue
            if not CJK.search(text):
                bad.append(f"{name} 行{r}: 详情缺中文")
            elif not EN_SENTENCE.search(text):
                bad.append(f"{name} 行{r}: 详情缺英文译文（机翻失败；清 _translate_cache.json 后重跑）")
            elif "官网未提供独立产品详情" not in text and "No standalone product description" not in text:
                hit += 1
    coverage = hit / total if total else 1.0
    if bad:
        return Gate("product_detail", "error", False, f"{len(bad)} 行产品详情不完整", bad)
    ok = (not empty) and coverage >= ctx.get("detail_min", 0.30)
    detail = (f"中英双语 {total - empty}/{total} 行；官网独立详情覆盖 {coverage:.0%}"
              f"（其余为分类/系列名双语占位说明）")
    offenders = [f"空详情 {empty} 行"] if empty else []
    return Gate("product_detail", "warn", ok, detail, offenders)


def gate_product_image_link(ctx):
    """产品图列必须是真实本地相对路径；官网无图必须显式标注。"""
    import openpyxl
    bad, no_image = [], []
    total = linked = 0
    for fp in glob.glob(os.path.join(ctx["deliverable"], "*", "产品清单.xlsx")):
        name = os.path.basename(os.path.dirname(fp))
        home = os.path.dirname(fp)
        try:
            ws = openpyxl.load_workbook(fp, read_only=True)["产品清单"]
        except Exception as exc:
            bad.append(f"{name}:无法读取 {exc}")
            continue
        hdr = [c.value for c in next(ws.iter_rows(max_row=1))]
        col = hdr.index("图片（本地连接）") if "图片（本地连接）" in hdr else 7
        for r, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            raw = str(row[col] if col < len(row) and row[col] is not None else "").strip()
            total += 1
            text = norm_text(raw)
            if not text:
                bad.append(f"{name} 行{r}: 本地链接为空")
                continue
            if "本地抓取，待核验" in text:
                bad.append(f"{name} 行{r}: 仍是旧占位文字")
                continue
            if "官网未提供产品图" in text or "No product image available" in text:
                no_image.append(f"{name} 行{r}: 官网未提供产品图")
                continue
            refs = [x.strip().replace("\\", "/") for x in raw.splitlines() if x.strip()]
            missing = []
            for ref in refs:
                target = os.path.normpath(os.path.join(home, *ref.split("/")))
                if not os.path.isfile(target):
                    missing.append(ref)
            if missing:
                bad.append(f"{name} 行{r}: 本地图片不存在 {'、'.join(missing[:3])}")
            else:
                linked += 1
    if bad:
        return Gate("product_image_link", "error", False, f"{len(bad)} 行图片本地链接无效", bad)
    detail = f"本地链接 {linked}/{total} 行；官网无图 {len(no_image)} 行（已显式标注）"
    return Gate("product_image_link", "warn", not no_image, detail, no_image)


def gate_product_rows(ctx):
    import openpyxl
    bad = []
    for name, d in sorted(ctx["companies"].items()):
        fp = os.path.join(company_path(ctx["deliverable"], name), "产品清单.xlsx")
        if not os.path.isfile(fp):
            continue
        try:
            ws = openpyxl.load_workbook(fp, read_only=True)["产品清单"]
        except Exception as exc:
            bad.append(f"{name}: 无法读取 {exc}")
            continue
        hdr = [c.value for c in next(ws.iter_rows(max_row=1))]
        col = hdr.index("产品名称(中英文)") if "产品名称(中英文)" in hdr else 2
        got = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            cell = row[col] if col < len(row) else None
            if cell is None or not str(cell).strip():
                continue
            got.append(norm_text(str(cell).split(" / ")[0]))
        want = [norm_text(p) for p in (d.get("products") or [])]
        if got != want:
            bad.append(f"{name}: 表 {len(got)} 行 / raw {len(want)} 个产品")
    return Gate("product_rows", "error", not bad, f"{len(bad)} 家产品清单行与 raw 不一致", bad)


def gate_product_map(ctx):
    bad = []
    for name, d in ctx["companies"].items():
        prods = d.get("products") or []
        if not prods:
            continue
        pm = ((ctx["en"] or {}).get(name) or {}).get("产品英名") or {}
        hit = sum(1 for p in prods if p in pm)
        if hit / len(prods) < ctx["map_min"]:
            bad.append(f"{name} {hit}/{len(prods)}")
    return Gate("product_map", "error", not bad,
                f"{len(bad)} 家产品英名覆盖率低于 {ctx['map_min']:.0%}", bad)


def gate_docx_sync(ctx):
    bad = []
    for name in sorted(ctx["companies"]):
        hits = glob.glob(os.path.join(ctx["deliverable"], name, "5.企业介绍", "*.docx"))
        if not hits:
            continue
        want = ((ctx["en"] or {}).get(name) or {}).get("英文简介") or []
        ps = docx_paragraphs(hits[0])
        if len(ps) < 4:
            bad.append(f"{name}: 段落不足")
            continue
        got = split_docx(ps)[2]
        if [t.strip() for t in want] != got:
            bad.append(f"{name}: docx {len(got)} 段 / 源 {len(want)} 段")
    return Gate("docx_sync", "error", not bad,
                f"{len(bad)} 家 docx 英文段与定稿不一致（改了数据未重渲染？）", bad)


def gate_docx_source(ctx):
    bad = []
    for name, d in sorted(ctx["companies"].items()):
        hits = glob.glob(os.path.join(company_path(ctx["deliverable"], name), "5.企业介绍", "*.docx"))
        if not hits:
            continue
        src = d.get("intro_paragraphs") or []
        if not src:
            src = [(d.get("home_text") or "")[:500] or "官网公开信息不足，待补充。"]
        want = [norm_text(t) for t in src[:3]]
        try:
            ps = docx_paragraphs(hits[0])
        except Exception as exc:
            bad.append(f"{name}: 无法读取 docx {exc}")
            continue
        got = [norm_text(t) for t in ps[1:] if CJK.search(t)]
        if got != want:
            if len(got) != len(want):
                why = f"docx 中文 {len(got)} 段 / raw {len(want)} 段"
            else:
                first = next((i for i, (g, w) in enumerate(zip(got, want), 1) if g != w), 1)
                why = f"第 {first} 段内容与 raw 不一致"
            bad.append(f"{name}: {why}")
    return Gate("docx_source", "error", not bad, f"{len(bad)} 家 docx 中文段与 raw 不一致", bad)


OK_VERDICTS = {"符合", "正确", "通过", "是", "ok", "pass", "yes"}
BAD_VERDICTS = {"不符", "不匹配", "错误", "否", "错", "no", "fail", "bad"}


def verdict_kind(value):
    s = str(value or "").strip().lower()
    if not s or s in ("待核对", "待定", "未核对", "pending"):
        return "pending"
    if s in OK_VERDICTS:
        return "ok"
    if s in BAD_VERDICTS or s.startswith("不符") or s.startswith("不匹配"):
        return "bad"
    return "pending"


def gate_visual_review(ctx):
    path = ctx.get("visual") or ""
    if not path or not os.path.isfile(path):
        level = "error" if ctx.get("require_visual") else "warn"
        return Gate("visual_review", level, False,
                    f"未找到视觉核对结论：{path or '<未指定>'}（先跑 visual_review.py）")
    try:
        data = read_json(path)
    except Exception as exc:
        return Gate("visual_review", "error", False, f"视觉核对文件无法解析：{exc}")
    bad, pending = [], []
    for name, d in sorted(ctx["companies"].items()):
        entry = data.get(name) if isinstance(data.get(name), dict) else {}
        if not entry:
            pending.append(f"{name}: 未核对")
            continue
        cats = entry.get("类别") if isinstance(entry.get("类别"), dict) else {}
        imgs = entry.get("图片") if isinstance(entry.get("图片"), dict) else {}
        for cat, folder in DIRS2.items():
            for it in d.get(cat) or []:
                key = f"{folder}/{it.get('file', '')}"
                per_image = imgs.get(key) if isinstance(imgs.get(key), dict) else {}
                kind = verdict_kind(per_image.get("结论"))
                if kind == "pending":
                    kind = verdict_kind((cats.get(folder) or {}).get("结论"))
                if kind == "bad":
                    bad.append(f"{name} {key}: 视觉核对不符")
                elif kind == "pending":
                    pending.append(f"{name} {key}: 待核对")
        for field in ("文档", "产品清单"):
            node = entry.get(field) if isinstance(entry.get(field), dict) else {}
            kind = verdict_kind(node.get("结论"))
            if kind == "bad":
                bad.append(f"{name} {field}: 视觉核对不符")
            elif kind == "pending":
                pending.append(f"{name} {field}: 待核对")
    if bad:
        return Gate("visual_review", "error", False,
                    f"{len(bad)} 条视觉核对不符、{len(pending)} 条待核对", bad + pending)
    if pending:
        level = "error" if ctx.get("require_visual") else "warn"
        return Gate("visual_review", level, False, f"{len(pending)} 条待 Codex 看图核对（未回写结论）", pending)
    return Gate("visual_review", "warn", True, f"{len(ctx['companies'])} 家视觉核对完成")


def gate_noise(ctx):
    bad = []
    for name in sorted(ctx["companies"]):
        hits = glob.glob(os.path.join(ctx["deliverable"], name, "5.企业介绍", "*.docx"))
        if not hits:
            continue
        ps = docx_paragraphs(hits[0])
        if len(ps) < 4:
            continue
        for i, t in enumerate(ps[1:], start=2):
            m = NOISE.search(t) or PLACEHOLDER.search(t)
            if m:
                bad.append(f"{name} 段{i}: {m.group(0)}")
    return Gate("noise", "error", not bad, f"{len(bad)} 处正文噪声", bad)


def gate_summary(ctx):
    import openpyxl
    if not os.path.isfile(ctx["summary"]):
        return Gate("summary", "error", False, f"汇总表不存在：{ctx['summary']}")
    ws = openpyxl.load_workbook(ctx["summary"], read_only=True).active
    hdr = [c.value for c in next(ws.iter_rows(max_row=1))]
    want = ["企业名称", "官网", "置信度", "处理地", "产业", "规模", "主要业务", "地址",
            "核心产品", "图片总数", "logo数", "产品图数", "资质图数", "工厂图数",
            "产品名数", "简介段数", "英文名称", "官网链接", "已生成文件夹"]
    bad = []
    if hdr != want:
        bad.append(f"表头不符：{hdr[:5]}…")
    rows = ws.max_row - 1
    exp = ctx["expected"]
    if exp and rows != exp:
        bad.append(f"{rows} 行 != 期望 {exp}")
    return Gate("summary", "error", not bad,
                "; ".join(bad) or f"{rows} 行 × {len(hdr)} 列", bad)


GATES = [
    gate_site_discovery,
    gate_coverage, gate_structure, gate_images, gate_image_required, gate_image_provenance,
    gate_en_entry, gate_en_ascii, gate_en_pinyin,
    gate_product_en, gate_product_detail, gate_product_image_link, gate_product_rows, gate_product_map,
    gate_docx_sync, gate_docx_source,
    gate_noise, gate_summary, gate_visual_review,
]


# ---------------------------------------------------------------------- 入口

def main():
    ap = argparse.ArgumentParser(description="企业官网资料包质量门禁")
    ap.add_argument("--config", default="", help="pipeline.json（路径与阈值）")
    ap.add_argument("--deliverable", default="", help="交付目录")
    ap.add_argument("--raw", default="", help="档案目录")
    ap.add_argument("--en", default="", help="定稿英文 json")
    ap.add_argument("--summary", default="", help="汇总 xlsx")
    ap.add_argument("--visual", default="", help="视觉核对.json，默认取 deliverable 同级")
    ap.add_argument("--require-visual", action="store_true", help="视觉核对未完成按 error 处理")
    ap.add_argument("--expected", type=int, default=0, help="期望企业数，0=不校验")
    ap.add_argument("--only", default="", help="只跑指定门禁，逗号分隔")
    ap.add_argument("--json", dest="json_out", default="", help="把结果写 json")
    ap.add_argument("--ascii-max", type=float, default=None, help="英文段非ASCII占比上限")
    ap.add_argument("--pinyin-max", type=float, default=None, help="疑似拼音占比告警线")
    ap.add_argument("--map-min", type=float, default=None, help="产品英名覆盖率下限")
    ap.add_argument("--detail-min", type=float, default=None, help="官网独立产品详情覆盖率告警线")
    a = ap.parse_args()

    cfg = {}
    path = expand(a.config)
    if not path:
        for cand in ("pipeline.json", os.path.join(os.getcwd(), "pipeline.json")):
            if os.path.isfile(cand):
                path = cand
                break
    if path and os.path.isfile(path):
        cfg = read_json(path)
        print(f"配置：{path}")
    elif a.config:
        raise SystemExit(f"找不到配置文件：{a.config}")

    def pick(name, cli):
        return expand(cli or cfg.get(name) or "")

    deliverable_path = pick("deliverable", a.deliverable)
    visual_path = pick("visual_review", a.visual)
    if not visual_path and deliverable_path:
        visual_path = os.path.join(os.path.dirname(os.path.abspath(deliverable_path)), "视觉核对.json")
    ctx = {
        "deliverable": deliverable_path,
        "raw": pick("raw", a.raw),
        "en_path": pick("en_final", a.en),
        "summary": pick("summary", a.summary),
        "visual": visual_path,
        "require_visual": bool(a.require_visual or cfg.get("require_visual")),
        "expected": a.expected or int(cfg.get("expected_companies") or 0),
        "ascii_max": (a.ascii_max if a.ascii_max is not None
                      else float(cfg.get("ascii_max", 0.02))),
        "pinyin_max": (a.pinyin_max if a.pinyin_max is not None
                       else float(cfg.get("pinyin_max", 0.60))),
        "map_min": (a.map_min if a.map_min is not None
                    else float(cfg.get("map_min", 0.80))),
        "detail_min": (a.detail_min if a.detail_min is not None
                       else float(cfg.get("detail_min", 0.30))),
    }
    for key in ("deliverable", "raw", "en_path", "summary"):
        if not ctx[key]:
            label = "en" if key == "en_path" else key.replace("_", "-")
            raise SystemExit(f"缺少参数：--{label}")
    if not os.path.isdir(ctx["raw"]):
        raise SystemExit(f"档案目录不存在：{ctx['raw']}")
    if not os.path.isdir(ctx["deliverable"]):
        raise SystemExit(f"交付目录不存在：{ctx['deliverable']}")

    ctx["companies"] = load_company_files(ctx["raw"])
    ctx["en"] = read_json(ctx["en_path"]) if ctx["en_path"] and os.path.isfile(ctx["en_path"]) else {}
    ctx["summary_names"] = None
    try:
        import openpyxl
        ws = openpyxl.load_workbook(ctx["summary"], read_only=True).active
        ctx["summary_rows"] = ws.max_row - 1
        ctx["summary_names"] = [
            str(row[0]).strip()
            for row in ws.iter_rows(min_row=2, values_only=True)
            if row and row[0] not in (None, "")
        ]
    except Exception:
        ctx["summary_rows"] = None

    only = {s.strip() for s in a.only.split(",") if s.strip()}
    known = {g.__name__.replace("gate_", "") for g in GATES}
    unknown = only - known
    if unknown:
        raise SystemExit(f"未知门禁：{', '.join(sorted(unknown))}；可选：{', '.join(sorted(known))}")
    todo = [g for g in GATES if not only or g.__name__.replace("gate_", "") in only]

    print(f"档案 {len(ctx['companies'])} 家 · 英文 {len(ctx['en'])} 家 · 交付 {ctx['deliverable']}\n")
    results, failed = [], 0
    for fn in todo:
        g = fn(ctx)
        results.append(g)
        mark = "PASS" if g.ok else ("FAIL" if g.level == "error" else "WARN")
        if not g.ok and g.level == "error":
            failed += 1
        print(f"[{mark}] {g.id:<12} {g.detail}")
        for o in g.offenders[:10]:
            print(f"         - {o}")
        if len(g.offenders) > 10:
            print(f"         … 另有 {len(g.offenders) - 10} 条")

    print(f"\n{'门禁全部通过' if not failed else f'{failed} 项 error 级门禁未通过'}")
    if a.json_out:
        os.makedirs(os.path.dirname(expand(a.json_out)) or ".", exist_ok=True)
        with open(expand(a.json_out), "w", encoding="utf-8") as fh:
            json.dump({"passed": not failed,
                       "gates": [g.as_dict() for g in results]}, fh,
                      ensure_ascii=False, indent=1)
        print(f"结果已写：{a.json_out}")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())