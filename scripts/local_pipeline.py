#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""本地企业官网资料包流水线。

输入一份 Excel（至少包含企业名称列），在本机完成：
官网发现 -> 页面抓取 -> 图片归档 -> 中文简介与产品提取 -> docx/xlsx 渲染 -> 门禁。

默认只用 Python 标准库和 openpyxl/python-docx，不依赖 SSH、opencode 或远端 worker。
Playwright 默认 auto：静态抓取为主，页面过薄或疑似 JS 渲染时自动用 Playwright 重抓。
"""
from __future__ import annotations

import argparse
import collections
import datetime as _dt
import gzip
import hashlib
import html as _html
import json
import mimetypes
import os
import re
import shutil
import ssl
import subprocess
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
import zlib
from html.parser import HTMLParser
from pathlib import Path

from domain_rules import is_directory_host, registrable

try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass

UA = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 CodexLocalPipeline/1.0"
IMAGE_DIRS = {
    "factory": "1.企业工厂图",
    "product": "2.企业产品图",
    "logo": "3.企业logo",
    "cert": "4.资质证书",
}
SUMMARY_HEADERS = [
    "企业名称", "官网", "置信度", "处理地", "产业", "规模", "主要业务", "地址",
    "核心产品", "图片总数", "logo数", "产品图数", "资质图数", "工厂图数",
    "产品名数", "简介段数", "英文名称", "官网链接", "已生成文件夹",
]
PRODUCT_COLUMNS = [
    "子品类", "排名", "产品名称(中英文)", "品牌", "价格人民币", "价格美金",
    "产品详情", "图片（本地连接）",
]
DETAIL_TAIL = re.compile(r"(?:查看详情|点击查看|了解更多|更多详情|立即咨询|马上咨询)\s*$")
NOISE = re.compile(
    r"首页[»>]|您当前的位置|欢迎光临|信息纠错|客服中心|会员级别|顺企|友情链接|"
    r"荟萃网库|未经核实|询盘|贸易通|立即注册|营业执照号码|技术支持|热门关键词|"
    r"Powered by|Toggle navigation|Jump to main|Sign in|Read ?more|Copyright|"
    r"查看详情|点击查看|了解更多|更多详情|立即咨询|马上咨询|"
    r"ICP备|备案号|版权所有|联系电话|联系方式",
    re.I,
)
SKIP_TAGS = {"script", "style", "noscript", "svg", "template", "iframe"}
SEARCH_BLOCK = re.compile(
    r"bing\.com|duckduckgo\.com|baidu\.com|so\.com|sogou\.com|google\.|"
    r"zhihu\.com|weibo\.com|douyin\.com|xiaohongshu\.com|qcc\.com|tianyancha\.com|"
    r"1688\.com|alibaba\.com|made-in-china\.com|yellowpages|黄页",
    re.I,
)
# 国内可达性优先：sogou / 360 会把真实结果域名暴露在页面 HTML 文本里，
# baidu 对无 Cookie 请求多数返回空页，duckduckgo 在部分网络下会被重置连接。
SEARCH_ENGINES = (
    ("sogou", "https://www.sogou.com/web?query={q}"),
    ("so360", "https://www.so.com/s?q={q}"),
    ("bing", "https://cn.bing.com/search?q={q}"),
    ("duckduckgo", "https://html.duckduckgo.com/html/?q={q}"),
)
# baidu 对无 Cookie 的 urllib 请求常返回空页；Playwright 可用时用它补搜。
PW_SEARCH_ENGINES = (
    ("baidu", "https://www.baidu.com/s?wd={q}"),
    ("bing", "https://cn.bing.com/search?q={q}"),
)
DOMAIN_RE = re.compile(
    r"(?<![0-9a-z@._-])((?:[a-z0-9](?:[a-z0-9-]{0,61}[a-z0-9])?\.)+"
    r"(?:com\.cn|net\.cn|org\.cn|gov\.cn|edu\.cn|com|cn|net|org|cc|top|vip|"
    r"tech|site|shop|info|biz|xin|store|online|ltd|group))",
    re.I,
)
JUNK_DOMAIN = re.compile(
    r"sogou|sogo\.|sgconst|so\.com|360\.|360tres|360sou|360kan|baidu|bing|microsoft|msn\.|"
    r"weixin|weibo|zhihu|douyin|xiaohongshu|eastmoney|xueqiu|zhipin|"
    r"qq\.com|live\.com|windowslive|azureedge|akamai|cloudfront|"
    r"qcc\.|tianyancha|maigoo|co188|10jqka|hao123|"
    r"google|gstatic|bdstatic|bcebos|qhimg|qpic|qqbrowser|wappass|gtimg|alicdn|"
    r"sinajs|tanx|mmstat|mediav|qhupdate|"
    r"window|console|document|prototype|undefined|localhost|w3\.org|json\.org|"
    r"crockford|github\.com|solib|"
    r"^[a-z0-9]\.(?:top|com|cn|net|org|info|biz)$",
    re.I,
)
def harvest_domains(html):
    """从搜索结果 HTML 里挖出候选域名及出现次数。"""
    counts = collections.Counter()
    for m in DOMAIN_RE.finditer(html or ""):
        d = m.group(1).lower().strip(".")
        sld = d.split(".")[0]
        if len(sld) < 3 or JUNK_DOMAIN.search(d) or is_directory_host(d):
            continue
        reg = registrable(d)
        if reg and not JUNK_DOMAIN.search(reg) and not is_directory_host(reg):
            counts[reg] += 1
    return counts


def alt_scheme(url):
    if url.startswith("https://"):
        return "http://" + url[len("https://"):]
    if url.startswith("http://"):
        return "https://" + url[len("http://"):]
    return ""


def fetch_site(url, args):
    """抓首页；https 不通时自动回退 http（很多国内企业站只开 http）。"""
    page, r = fetch_page(url, args)
    if page:
        return page, r
    alt = alt_scheme(url)
    if alt:
        page2, r2 = fetch_page(alt, args)
        if page2:
            return page2, r2
    return page, r
CJK_RE = re.compile(r"[\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff]")
TRANSLATE_URL = "https://api.mymemory.translated.net/get"
TRANSLATE_CHUNK = 450
# MyMemory 匿名查询约 500 字节上限；中文 1 字 ~3 字节，必须按字节而非字符分片。
TRANSLATE_CHUNK_BYTES = 400


def log(*parts):
    print(*parts, flush=True)


def clean_inline(s):
    return re.sub(r"\s+", " ", _html.unescape(str(s or ""))).strip()


def clean_lines(s):
    """保留换行/制表符的清洗：用于表格类资料，避免行列边界丢失。"""
    s = _html.unescape(str(s or "")).replace("\r\n", "\n").replace("\r", "\n")
    lines = []
    for line in s.split("\n"):
        line = re.sub(r"[ \u00a0\u3000]+", " ", line).strip()
        if line:
            lines.append(line)
    return "\n".join(lines)


def safe_filename(s, fallback="company"):
    s = clean_inline(s)
    s = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", s)
    s = s.strip(" .")
    return s[:120] or fallback


def normalize_key(s):
    return re.sub(r"[\s\W_]+", "", clean_inline(s), flags=re.UNICODE).lower()


def core_company_name(name):
    s = clean_inline(name)
    s = re.sub(r"（[^）]*(?:原|备注|变更)[^）]*）|\([^)]*(?:原|备注|变更)[^)]*\)", "", s)
    s = re.sub(r"^(?:四川省?|成都市?|德阳市?|绵阳市?|宜宾市?|泸州市?|乐山市?|眉山市?|资阳市?|遂宁市?|内江市?|南充市?|达州市?|广安市?|巴中市?|雅安市?|攀枝花市?)+", "", s)
    s = re.sub(r"(?:股份有限公司|有限责任公司|有限公司|集团公司|集团|公司|工厂|厂|研究院|研究所)$", "", s)
    return s or s.strip("有限公司") or clean_inline(name)


def norm_url(url):
    url = clean_inline(url)
    if not url:
        return ""
    if not re.match(r"^https?://", url, re.I):
        url = "https://" + url
    return url


def host_of(url):
    try:
        return urllib.parse.urlparse(url).netloc.lower().split("@")[-1].split(":")[0]
    except Exception:
        return ""


def same_site(a, b):
    ha, hb = host_of(a), host_of(b)
    return bool(ha and hb and (ha == hb or ha.endswith("." + hb) or hb.endswith("." + ha)))


def url_category(url, text=""):
    s = (clean_inline(url) + " " + clean_inline(text)).lower()
    checks = [
        ("cert", ("证书", "资质", "荣誉", "认证", "certificate", "honor", "award", "cert")),
        ("product", ("产品", "商品", "服务", "解决方案", "业务", "product", "service", "solution")),
        ("factory", ("工厂", "车间", "厂房", "设备", "生产", "基地", "factory", "workshop", "equipment", "plant")),
        ("about", ("关于", "简介", "公司介绍", "企业介绍", "profile", "about", "company")),
    ]
    for cat, words in checks:
        if any(w in s for w in words):
            return cat
    return ""


IMG_ALT_NOISE = (
    "导航", "收起", "展开", "联系", "热线", "电话",
    "二维码", "扫一扫", "关注", "分享", "客服", "咨询",
    "返回", "顶部", "底部", "菜单", "首页", "微信", "微博",
    "公众号", "喜报", "揭牌", "有限公司",
    "menu", "nav", "close", "more", "back", "qr", "wechat", "weixin", "icon",
)


CONTACT_ALT_RE = re.compile(
    r"[\w.+-]+@[\w-]+\.[\w.-]+"                      # 邮箱
    r"|\+\d[\d\s\-()]{6,}"                           # 国际电话 +86 ...
    r"|(?<!\d)\d{7,}(?!\d)"                            # 7 位以上号码
    r"|(?:省|自治区).{0,10}(?:市|区|县|镇|路|街|工业园区|工业园|大厦)"
    r"|(?:address|tel|phone|e-?mail)\b", re.I)


def is_contact_alt(alt):
    """判断图片 alt 是否为联系方式（电话/邮箱/地址），用于剔除网站图标。"""
    return bool(alt and CONTACT_ALT_RE.search(alt))


def image_category(img, page_cat):
    alt = clean_inline(img.get("alt", ""))
    s = (clean_inline(img.get("src", "")) + " " + alt).lower()
    if img.get("bg"):
        # 背景横幅：仅当出现"关于/公司/厂区"线索时才作为工厂/厂区候选，其余丢弃。
        if any(w in s for w in ("about", "company", "factory", "plant", "profile",
                                "工厂", "厂区", "车间", "基地", "公司", "简介")):
            return "factory"
        return ""
    if any(w in s for w in ("logo", "标志", "徽标", "商标", "brand-logo")):
        return "logo"
    # 位于指向首页的链接内、且是该链接第一张图：通常是站头 logo。
    # 部分模板 logo 无 alt 或 alt 恰为公司名，必须在 alt 噪声过滤前判定。
    if img.get("home_link"):
        return "logo"
    if alt and is_contact_alt(alt):
        return ""
    if alt and any(w in alt.lower() for w in IMG_ALT_NOISE):
        return ""
    if any(w in s for w in ("cert", "证书", "资质", "荣誉", "认证", "award")):
        return "cert"
    if any(w in s for w in ("product", "产品", "商品", "goods", "pro_")):
        return "product"
    if any(w in s for w in ("factory", "工厂", "车间", "厂房", "设备", "生产", "workshop", "plant")):
        return "factory"
    return page_cat if page_cat in ("factory", "product", "cert", "logo") else ""


def is_home_href(href):
    """判断链接是否指向站点首页（用于识别站头 logo 所在的链接）。

    兼容相对写法（/、./index.html）与同站绝对写法（http://host/、https://host/index.html）。
    """
    h = clean_inline(href).split("#")[0].split("?")[0].strip().lower()
    if h in ("/", "./", "./index.html", "index.html", "/index.html"):
        return True
    if "://" in h:
        rest = h.split("://", 1)[1]
        path = rest[rest.find("/"):] if "/" in rest else "/"
        return path in ("", "/", "/index.html")
    return False


class PageParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.title = []
        self._title = None
        self.text = []
        self.links = []
        self.images = []
        self.headings = []
        self.lists = []
        self._skip = 0
        self._a = None
        self._h = None
        self._li = None

    def handle_starttag(self, tag, attrs):
        tag = tag.lower()
        d = {k.lower(): v for k, v in attrs}
        if tag in SKIP_TAGS:
            self._skip += 1
            return
        if self._skip:
            return
        if tag == "title":
            self._title = []
        if tag == "img":
            src = d.get("src") or d.get("data-src") or d.get("data-original") or d.get("data-lazy-src") or ""
            if not src and d.get("srcset"):
                src = d["srcset"].split(",")[0].strip().split(" ")[0]
            home_link = False
            if self._a is not None:
                home_link = (self._a.get("img_count", 0) == 0
                             and is_home_href(self._a.get("href", "")))
                self._a["img_count"] = self._a.get("img_count", 0) + 1
            self.images.append({
                "src": src, "alt": d.get("alt", ""),
                "width": d.get("width", ""), "height": d.get("height", ""),
                "home_link": home_link,
            })
        if tag == "a":
            self._a = {"href": d.get("href", ""), "text": [], "img_count": 0}
        elif tag in ("h1", "h2", "h3", "h4", "h5"):
            self._h = {"tag": tag, "text": []}
        elif tag == "li":
            self._li = []
        elif tag == "br" and self._li is not None:
            self._li.append(" ")

    def handle_endtag(self, tag):
        tag = tag.lower()
        if tag in SKIP_TAGS:
            if self._skip:
                self._skip -= 1
            return
        if self._skip:
            return
        if tag == "a" and self._a is not None:
            self.links.append({"href": self._a["href"], "text": clean_inline(" ".join(self._a["text"]))})
            self._a = None
        elif tag in ("h1", "h2", "h3", "h4", "h5") and self._h is not None:
            t = clean_inline(" ".join(self._h["text"]))
            if t:
                self.headings.append(t)
            self._h = None
        elif tag == "li" and self._li is not None:
            t = clean_inline(" ".join(self._li))
            if t:
                self.lists.append(t)
            self._li = None
        elif tag == "title" and self._title is not None:
            t = clean_inline(" ".join(self._title))
            if t:
                self.title.append(t)
            self._title = None

    def handle_data(self, data):
        if self._skip:
            return
        s = clean_inline(data)
        if not s:
            return
        self.text.append(s)
        if self._title is not None:
            self._title.append(s)
        if self._a is not None:
            self._a["text"].append(s)
        if self._h is not None:
            self._h["text"].append(s)
        if self._li is not None:
            self._li.append(s)


def decompress_body(body, headers):
    """按 Content-Encoding 解压响应体；服务端强制 gzip 时静态抓取不会自动解压。"""
    enc = ""
    try:
        enc = (headers.get("Content-Encoding") or "").strip().lower()
    except Exception:
        enc = ""
    if not enc or enc == "identity" or not body:
        return body
    try:
        if "gzip" in enc:
            return gzip.decompress(body)
        if "deflate" in enc:
            try:
                return zlib.decompress(body)
            except Exception:
                return zlib.decompress(body, -zlib.MAX_WBITS)
    except Exception:
        return body
    return body


def sanitize_text(text):
    """剔除 XML 控制字符，避免写 docx 时抛 XML compatible 错误。"""
    if not text:
        return text
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)
    return re.sub(r"[\ud800-\udfff]", "", text)


def decode_body(body, headers):
    body = decompress_body(body, headers)
    charset = ""
    try:
        charset = headers.get_content_charset() or ""
    except Exception:
        pass
    if not charset:
        m = re.search(br"charset\s*=\s*[\"']?([A-Za-z0-9._-]+)", body[:8192], re.I)
        if m:
            charset = m.group(1).decode("ascii", "ignore")
    for enc in (charset, "utf-8", "gb18030", "big5"):
        if not enc:
            continue
        try:
            return sanitize_text(body.decode(enc))
        except Exception:
            continue
    return sanitize_text(body.decode("utf-8", "replace"))


def build_opener(proxy="", insecure=False):
    handlers = []
    if proxy:
        handlers.append(urllib.request.ProxyHandler({"http": proxy, "https": proxy}))
    if insecure:
        handlers.append(urllib.request.HTTPSHandler(context=ssl._create_unverified_context()))
    return urllib.request.build_opener(*handlers)


def fetch(url, args, binary=False, referer=""):
    url = norm_url(url)
    headers = {
        "User-Agent": UA,
        "Accept": "image/avif,image/webp,image/apng,image/svg+xml,image/*,*/*;q=0.8" if binary
                  else "text/html,application/xhtml+xml,application/xml,*/*;q=0.8",
        "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.6",
    }
    if referer:
        headers["Referer"] = referer
    req = urllib.request.Request(url, headers=headers)
    try:
        with build_opener(args.proxy, args.insecure).open(req, timeout=args.timeout) as resp:
            body = resp.read(args.max_bytes)
            return {
                "url": resp.geturl(), "headers": resp.headers, "body": body,
                "binary": binary, "status": getattr(resp, "status", ""), "error": "",
            }
    except Exception as exc:
        return {"url": url, "headers": {}, "body": b"", "binary": binary,
                "status": "", "error": str(exc)}


def _byte_clip(text, byte_limit):
    """截取不超过 byte_limit 个 UTF-8 字节的前缀。"""
    used, end = 0, 0
    for idx, ch in enumerate(text):
        n = len(ch.encode("utf-8"))
        if used + n > byte_limit:
            break
        used += n
        end = idx + 1
    return text[:end]


def chunk_text(text, limit=TRANSLATE_CHUNK, byte_limit=TRANSLATE_CHUNK_BYTES):
    """按句子边界切片，每片同时不超过 limit 字符和 byte_limit 个 UTF-8 字节。"""
    text = clean_inline(text)
    if not text:
        return []
    out, buf = [], ""
    for sent in re.split(r"(?<=[。！？；!?;])", text):
        sent = sent.strip()
        if not sent:
            continue
        while len(sent) > limit or len(sent.encode("utf-8")) > byte_limit:
            clip = _byte_clip(sent, byte_limit)
            if len(clip) > limit:
                clip = clip[:limit]
            if not clip:
                break
            if buf:
                out.append(buf)
                buf = ""
            out.append(clip)
            sent = sent[len(clip):].strip()
        if not sent:
            continue
        joined = buf + sent
        if buf and (len(joined) > limit or len(joined.encode("utf-8")) > byte_limit):
            out.append(buf)
            buf = sent
        else:
            buf = joined
    if buf:
        out.append(buf)
    return out


_RATE_LIMIT_UNTIL = 0.0
_RATE_LIMIT_STREAK = 0
_RATE_LIMIT_OPEN_UNTIL = 0.0
_RATE_LIMIT_TRIP = 3


def translate_chunk(text, args):
    """调用 MyMemory 免费接口翻译一段中文；失败返回空串。

    429 限流时做冷却退避；连续多次整条失败后只熔断一个时间窗，窗口过后自动
    放行探测，避免要么空转十几分钟、要么把整批英文永久丢掉。
    """
    global _RATE_LIMIT_UNTIL, _RATE_LIMIT_STREAK, _RATE_LIMIT_OPEN_UNTIL
    if _RATE_LIMIT_STREAK >= _RATE_LIMIT_TRIP:
        if time.time() < _RATE_LIMIT_OPEN_UNTIL:
            return ""
        _RATE_LIMIT_STREAK = 0      # 冷却窗口已过，重新探测
    query = urllib.parse.urlencode({"q": text, "langpair": "zh-CN|en"})
    email = clean_inline(getattr(args, "translate_email", ""))
    if email:
        query += "&de=" + urllib.parse.quote(email)
    pause = _RATE_LIMIT_UNTIL - time.time()
    if pause > 0:
        time.sleep(min(pause, 10))
    for wait in (0, 3.0, 8.0):
        if wait:
            time.sleep(wait)
        r = fetch(f"{TRANSLATE_URL}?{query}", args)
        err = str(r.get("error") or "")
        if "429" in err:
            # 限流是接口级状态，重试本条无意义：记冷却时间后直接判失败。
            _RATE_LIMIT_UNTIL = time.time() + 10
            break
        if err or not r.get("body"):
            continue
        try:
            payload = json.loads(r["body"].decode("utf-8", "replace"))
        except Exception:
            continue
        got = clean_inline((payload.get("responseData") or {}).get("translatedText") or "")
        if got and "MYMEMORY WARNING" not in got.upper():
            _RATE_LIMIT_STREAK = 0
            _RATE_LIMIT_OPEN_UNTIL = 0.0
            return _html.unescape(got)
    _RATE_LIMIT_STREAK += 1
    if _RATE_LIMIT_STREAK >= _RATE_LIMIT_TRIP:
        _RATE_LIMIT_OPEN_UNTIL = time.time() + 60
    return ""


class Translator:
    """带磁盘缓存的中译英器，用于自动生成英文草稿。"""

    def __init__(self, args, cache_path=None):
        self.args = args
        self.cache_path = Path(cache_path) if cache_path else None
        self.cache = {}
        self.fail = 0
        if self.cache_path and self.cache_path.is_file():
            try:
                loaded = json.loads(self.cache_path.read_text(encoding="utf-8"))
                if isinstance(loaded, dict):
                    # 丢掉历史空结果：一次限流不应被永久固化在缓存里。
                    self.cache = {k: v for k, v in loaded.items() if v}
            except Exception:
                self.cache = {}

    def translate(self, text):
        text = clean_inline(text)
        if not text:
            return ""
        if not CJK_RE.search(text):
            return text
        if text in self.cache:
            return self.cache[text]
        parts = [translate_chunk(chunk, self.args) for chunk in chunk_text(text)]
        ok = [p for p in parts if p]
        out = clean_inline(" ".join(ok))
        delay = getattr(self.args, "translate_delay", 0)
        if len(ok) < len(parts):
            # 整条或部分失败：不写缓存，下轮重试，避免限流导致英文永久缺失。
            self.fail += 1
            if delay:
                time.sleep(delay)
            return out
        self.cache[text] = out
        if delay:
            time.sleep(delay)
        return out

    def save(self):
        if not self.cache_path:
            return
        try:
            self.cache_path.parent.mkdir(parents=True, exist_ok=True)
            self.cache_path.write_text(
                json.dumps(self.cache, ensure_ascii=False, indent=0), encoding="utf-8")
        except Exception:
            pass


def english_sentences(text):
    return [s.strip() for s in re.split(r"(?<=[.!?])\s+", clean_inline(text)) if s.strip()]


def split_balanced(sents, groups=3):
    """把句子按字符长度均衡分成 groups 段。"""
    total = sum(len(s) for s in sents)
    out, cur, acc = [], [], 0
    for s in sents:
        cur.append(s)
        acc += len(s)
        if len(out) < groups - 1 and acc >= total * (len(out) + 1) / groups:
            out.append(" ".join(cur).strip())
            cur, acc = [], 0
    if cur:
        out.append(" ".join(cur).strip())
    return [p for p in out if p]


def build_english_intro(intro_paragraphs, translator, fallback_text=""):
    """逐段翻译中文简介，保持原来的段落划分；不足三段时再按长度均衡拆分。"""
    translated = [translator.translate(p) for p in (intro_paragraphs or [])]
    translated = [t for t in translated if t and not CJK_RE.search(t)]
    if not translated and fallback_text:
        got = translator.translate(fallback_text)
        if got and not CJK_RE.search(got):
            translated = [got]
    if len(translated) >= 3:
        return translated[:3]
    sents = []
    for para in translated:
        sents += english_sentences(para)
    if len(sents) <= len(translated):
        return translated
    return split_balanced(sents, 3)


def derive_brand(en_name):
    """从英文企业名取品牌主体（去掉公司后缀）。"""
    s = clean_inline(en_name)
    s = re.sub(r"[,\s]+(?:Co\.?,?\s*)?(?:Ltd\.?|Limited|Inc\.?|LLC|Corp\.?|Corporation|Group|Holdings?)\b.*$",
               "", s, flags=re.I)
    s = re.sub(r"\b(?:Group|Holdings?|Company|Co\.?|Ltd\.?|Limited|Inc\.?|LLC|Corp\.?|Corporation)\b",
               " ", s, flags=re.I)
    return clean_inline(re.sub(r"\s{2,}", " ", s)).strip(" ,") or clean_inline(en_name)


def auto_english_entry(archive, translator):
    """用自动翻译生成英文层：英文名、英文标题、品牌、三段简介和产品英名。"""
    en_name = translator.translate(archive.get("名称", ""))
    en_paras = build_english_intro(
        archive.get("intro_paragraphs"), translator,
        fallback_text=clean_inline(archive.get("home_text", ""))[:400])
    products = archive.get("products") or []
    product_map = {}
    for product in products:
        eng = translator.translate(product)
        if eng and not CJK_RE.search(eng):
            product_map[product] = eng
    detail_map = {}
    details = archive.get("product_details") or {}
    detail_items = [(p, clean_inline(d)) for p, d in details.items() if clean_inline(d)]
    for product, detail in detail_items:
        eng = translator.translate(detail)
        if eng and not CJK_RE.search(eng):
            detail_map[product] = eng
    return {
        "英文名": en_name,
        "英文标题": en_name,
        "品牌": derive_brand(en_name),
        "英文简介": en_paras,
        "产品英名": product_map,
        "产品详情英": detail_map,
        "子品类": archive.get("产业", ""),
        "定稿": False,
        "自动翻译": True,
        "翻译引擎": "MyMemory",
        "翻译时间": _dt.datetime.now().isoformat(timespec="seconds"),
        "翻译失败": ((not en_name) or (not en_paras)
                     or (len(product_map) < len(products))
                     or (len(detail_map) < len(detail_items))),
        "备注": "英文为自动翻译草稿，待人工核校",
    }


def empty_english_entry(archive):
    return {
        "英文名": "", "英文标题": archive.get("名称", ""), "品牌": "", "英文简介": [],
        "产品英名": {}, "产品详情英": {}, "子品类": archive.get("产业", ""),
        "定稿": False, "自动翻译": False, "翻译失败": True,
        "备注": "未启用自动翻译；需补 --en 或明确接受中文版",
    }


def english_backlog(archives, en_data):
    """列出自动翻译失败或英文层缺口，供 Codex 生成 --en 补译稿。"""
    items = {}
    for archive in archives:
        name = archive.get("名称", "")
        entry = en_data.get(name) or {}
        products = archive.get("products") or []
        details = archive.get("product_details") or {}
        detail_source = {p: clean_inline(d) for p, d in details.items() if clean_inline(d)}
        product_en = entry.get("产品英名") if isinstance(entry.get("产品英名"), dict) else {}
        detail_en = entry.get("产品详情英") if isinstance(entry.get("产品详情英"), dict) else {}
        missing_products = [p for p in products if not clean_inline(product_en.get(p, ""))]
        missing_details = [p for p in detail_source if not clean_inline(detail_en.get(p, ""))]
        paras = entry.get("英文简介") if isinstance(entry.get("英文简介"), list) else []
        problems = []
        if entry.get("翻译失败") is True:
            problems.append("自动翻译未完成")
        if not clean_inline(entry.get("英文名", "")):
            problems.append("英文名为空")
        if len([p for p in paras if clean_inline(p)]) != 3:
            problems.append("英文简介不足三段")
        if missing_products:
            problems.append("产品英名缺 %d 个" % len(missing_products))
        if missing_details:
            problems.append("产品详情英缺 %d 个" % len(missing_details))
        if problems:
            items[name] = {
                "问题": problems,
                "建议": "基于 raw 中文事实生成 --en 兼容 JSON；保留 定稿:false、自动翻译:true、备注:待人工核校",
                "缺产品英名": missing_products,
                "缺产品详情英": missing_details,
            }
    return {
        "说明": "自动英文不完整。默认门禁禁止发布；Codex 必须用 --en 补稿，或由用户明确 --accept-no-english 接受中文版。",
        "企业": items,
    }


def playwright_status():
    """本地是否已安装 Playwright Python 包。"""
    try:
        import importlib.util
        return importlib.util.find_spec("playwright") is not None
    except Exception:
        return False


EDGE_REL = os.path.join("Microsoft", "Edge", "Application", "msedge.exe")
_BROWSER_CHOICE = None


def find_edge():
    """定位本机 Microsoft Edge；找不到返回空串。"""
    exe = shutil.which("msedge") or shutil.which("msedge.exe")
    if exe:
        return exe
    roots = [
        os.environ.get("ProgramFiles(x86)"),
        os.environ.get("ProgramFiles"),
        os.environ.get("LOCALAPPDATA"),
        os.environ.get("ProgramW6432"),
    ]
    for root in roots:
        if not root:
            continue
        cand = os.path.join(root, EDGE_REL)
        if os.path.isfile(cand):
            return cand
    return ""


def browser_launch_options():
    """按优先级给出可用的渲染内核：(标签, launch 参数)。

    内置 Chromium -> Edge 稳定通道 -> Edge 绝对路径。
    Edge 与 Chromium 同源，渲染能力一致，但系统自带、无需下载 100-200MB。
    """
    global _BROWSER_CHOICE
    if _BROWSER_CHOICE is not None:
        return [_BROWSER_CHOICE]
    options = [("chromium", {})]
    edge = find_edge()
    if edge:
        options.append(("msedge", {"channel": "msedge"}))
        options.append(("msedge-exe", {"executable_path": edge}))
    return options


def launch_rendered_browser(p):
    """启动第一个可用的浏览器内核，返回 (browser, 标签)。"""
    global _BROWSER_CHOICE
    last = None
    for label, kwargs in browser_launch_options():
        try:
            browser = p.chromium.launch(headless=True, **kwargs)
            _BROWSER_CHOICE = (label, kwargs)
            return browser, label
        except Exception as exc:
            last = exc
    raise RuntimeError("Chromium/Edge 均不可用：" + str(last))


def playwright_mode(args):
    mode = getattr(args, "playwright", "auto")
    if mode is True:
        return "on"
    if mode is False:
        return "off"
    mode = str(mode or "auto").lower()
    return mode if mode in ("auto", "on", "off") else "auto"


def playwright_fetch(url, args):
    """用 Chromium 渲染页面。auto 模式下仅在静态结果过薄时调用。"""
    try:
        from playwright.sync_api import sync_playwright
    except Exception as exc:
        return {"url": url, "headers": {}, "body": b"", "binary": False,
                "status": "", "error": f"未安装 Playwright：{exc}"}
    browser = None
    try:
        with sync_playwright() as p:
            browser, _label = launch_rendered_browser(p)
            page = browser.new_page(locale="zh-CN", user_agent=UA)
            page.goto(norm_url(url), wait_until="domcontentloaded",
                      timeout=max(5000, args.timeout * 1000))
            try:
                page.wait_for_load_state("networkidle", timeout=min(8000, max(2000, args.timeout * 500)))
            except Exception:
                pass
            page.wait_for_timeout(500)
            final_url = page.url
            body = page.content().encode("utf-8")
            return {"url": final_url, "headers": {}, "body": body, "binary": False,
                    "status": 200, "error": ""}
    except Exception as exc:
        return {"url": url, "headers": {}, "body": b"", "binary": False,
                "status": "", "error": str(exc)}
    finally:
        try:
            if browser is not None:
                browser.close()
        except Exception:
            pass


def parse_page(raw, fallback_url=""):
    """把一次抓取结果解析成统一 Page 结构；无内容返回 None。"""
    if not raw or raw.get("error") or not raw.get("body"):
        return None
    text = decode_body(raw["body"], raw["headers"])
    p = PageParser()
    try:
        p.feed(text)
    except Exception:
        pass
    page = {
        "url": raw.get("url") or fallback_url,
        "title": clean_inline(" ".join(p.title)),
        "text": clean_inline(" ".join(p.text)),
        "headings": p.headings,
        "lists": p.lists,
        "links": [],
        "images": [],
    }
    base = page["url"] or fallback_url
    for item in p.links:
        href = urllib.parse.urljoin(base, item["href"])
        if href.startswith(("http://", "https://")):
            page["links"].append({"href": href, "text": item["text"]})
    for item in p.images:
        src = urllib.parse.urljoin(base, item["src"].replace("\\", "/"))
        if src.startswith(("http://", "https://")):
            item["src"] = src
            page["images"].append(item)
    # CSS 背景图不会被 <img> 解析器捕获；首页横幅常以 background-image 方式出现，
    # 其中"关于/公司/厂区"横幅是工厂图的重要候选来源。
    seen_src = {item.get("src", "") for item in page["images"]}
    for m in re.finditer(
        r"background(?:-image)?\s*:\s*url\(\s*[\"']?([^\"')]+?)[\"']?\s*\)",
        text, re.I,
    ):
        raw_src = clean_inline(m.group(1)).replace("\\", "/")
        if not raw_src or raw_src.startswith("data:"):
            continue
        src = urllib.parse.urljoin(base, raw_src)
        if src.startswith(("http://", "https://")) and src not in seen_src:
            seen_src.add(src)
            page["images"].append({
                "src": src, "alt": "", "width": "", "height": "",
                "home_link": False, "bg": True,
            })
    return page


def page_is_thin(page):
    """静态 HTML 抓到的正文/链接/图片过少，判断为疑似 JS 渲染页。"""
    if not page:
        return True
    text_len = len(page.get("text") or "")
    return (text_len < 400 and len(page.get("links") or []) < 5
            and len(page.get("images") or []) < 3)


def norm_key(url):
    return norm_url(url).rstrip("/").lower()


def load_html_overrides(args):
    """载入 --html-dir 下 Codex 内置浏览器保存的离线 HTML。

    manifest.json 形如 {"https://a.com/": "a.html", ...}。
    """
    root = Path(getattr(args, "html_dir", "") or "").expanduser()
    if not str(getattr(args, "html_dir", "") or "").strip():
        return {}
    manifest = root / "manifest.json"
    if not manifest.is_file():
        log(f"[html-dir] 未找到 {manifest}，忽略离线页面")
        return {}
    try:
        mapping = json.loads(manifest.read_text(encoding="utf-8"))
    except Exception as exc:
        log(f"[html-dir] manifest.json 解析失败：{exc}")
        return {}
    out = {}
    for url, rel in (mapping or {}).items():
        f = (root / str(rel)).resolve()
        if f.is_file():
            out[norm_key(str(url))] = f
    log(f"[html-dir] 载入 {len(out)} 个离线页面快照")
    return out


def html_override(url, args):
    """离线快照优先于网络抓取；没有则返回 None。"""
    overrides = getattr(args, "html_overrides", None) or {}
    f = overrides.get(norm_key(url))
    if f is None:
        return None
    try:
        raw = {"url": url, "headers": {}, "body": f.read_bytes(), "binary": False,
               "status": 200, "error": "", "offline": True}
    except Exception as exc:
        log(f"[html-dir] 读取失败 {f}：{exc}")
        return None
    page = parse_page(raw, url)
    if page is None:
        return None
    log(f"[html-dir] 使用离线快照：{url}")
    return page, raw


def fetch_page(url, args):
    """抓页面并解析。

    默认 auto：先静态抓取；结果过薄或失败且本地有 Playwright 时自动重抓。
    on：强制 Playwright（失败回退静态）；off：只静态抓取。
    """
    override = html_override(url, args)
    if override is not None:
        return override
    mode = playwright_mode(args)
    if mode == "on":
        raw = playwright_fetch(url, args)
        page = parse_page(raw, url)
        if page is not None:
            return page, raw
        raw2 = fetch(url, args)
        page2 = parse_page(raw2, url)
        return page2, (raw2 if page2 is not None else raw)

    raw = fetch(url, args)
    page = parse_page(raw, url)
    if mode == "auto" and playwright_status() and page_is_thin(page):
        raw2 = playwright_fetch(url, args)
        page2 = parse_page(raw2, url)
        if page2 is not None:
            better = (page is None
                      or len(page2.get("text") or "") > len(page.get("text") or "") * 1.2
                      or len(page2.get("images") or []) > len(page.get("images") or [])
                      or len(page2.get("links") or []) > len(page.get("links") or []))
            if better:
                return page2, raw2
    return page, raw


def load_excel(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = [clean_inline(x) for x in rows[0]]
    def find_col(words):
        for i, h in enumerate(header):
            if any(w in h for w in words):
                return i
        return None
    name_col = find_col(("企业名称", "公司名称", "单位名称", "名称"))
    site_col = find_col(("官网", "网址", "网站", "官网地址"))
    has_header = name_col is not None
    if name_col is None:
        name_col = 0
    out = []
    for r in rows[1:] if has_header else rows:
        name = clean_inline(r[name_col] if name_col < len(r) else "")
        if not name:
            continue
        site = clean_inline(r[site_col] if site_col is not None and site_col < len(r) else "")
        out.append({"name": name, "website": site})
    wb.close()
    return out


def search_candidates(name, args):
    """多引擎挖掘并按“与公司名共现”加权排序候选官网域名。

    国内网络下 sogou/360 可达性最好；静态结果过少且本机有 Playwright 时，
    再用 baidu/bing 渲染页补搜。返回按可能性排序的候选 URL。
    """
    core = core_company_name(name)
    aliases = [name] if not core or core == name else [name, core]
    queries = []
    for alias in aliases:
        queries += [f'"{alias}" 官网', f'"{alias}" 官方网站']
    queries.append(f'"{name}" ICP备案')
    seen_q = set()
    queries = [q for q in queries if not (q in seen_q or seen_q.add(q))]

    counts = collections.Counter()
    ctx_counts = collections.Counter()
    order = {}
    full_key = normalize_key(name)
    core_key = normalize_key(core or name)

    def harvest(html, weight=1):
        hits = harvest_domains(html)
        for reg, cnt in hits.items():
            if reg not in order:
                order[reg] = len(order)
            counts[reg] += cnt * weight
        # 域名出现在公司名附近的上下文，权重远高于全页裸域名。
        for m in DOMAIN_RE.finditer(html or ""):
            reg = registrable(m.group(1).lower().strip("."))
            if not reg or JUNK_DOMAIN.search(reg) or is_directory_host(reg):
                continue
            window = normalize_key(html[max(0, m.start() - 220):m.start() + 220])
            if (full_key and full_key in window) or (core_key and core_key in window):
                ctx_counts[reg] += 1
                if reg not in order:
                    order[reg] = len(order)

    for query in queries[:6]:
        for _engine, tpl in SEARCH_ENGINES:
            r = fetch(tpl.format(q=urllib.parse.quote(query)), args)
            if r.get("error") or not r.get("body"):
                continue
            html = decode_body(r["body"], r["headers"])
            harvest(html)
            for m in re.finditer(r"https?://([^/\"'\s<>\\]+)", html):
                host = m.group(1).split("@")[-1].split(":")[0].lower()
                reg = registrable(host)
                if (reg and len(reg.split(".")[0]) >= 3
                        and not JUNK_DOMAIN.search(reg) and not is_directory_host(reg)):
                    if reg not in order:
                        order[reg] = len(order)
                    counts[reg] += 1

    # 静态搜索覆盖不足时，用 Playwright 补搜（baidu 对无 Cookie 请求基本不可用）。
    if len(counts) < 3 and playwright_status() and playwright_mode(args) != "off":
        for query in queries[:2]:
            for _engine, tpl in PW_SEARCH_ENGINES:
                r = playwright_fetch(tpl.format(q=urllib.parse.quote(query)), args)
                if r.get("error") or not r.get("body"):
                    continue
                harvest(decode_body(r["body"], r["headers"]), weight=2)

    ranked = sorted(
        counts,
        key=lambda d: (-(counts[d] + ctx_counts[d] * 5), order.get(d, 0), d),
    )
    return [f"https://{d}/" for d in ranked[:max(1, args.search_candidates)]]


def site_score(name, page):
    """给候选首页打分：公司名命中标题/正文是主证据，联系方式/备案是补强。

    黄页/工商/名录/B2B 聚合站会把企业名写进列表页，仅靠命中无法与官网区分，
    因此这类域名直接判 0 分，不参与候选排序。
    """
    if not page:
        return 0
    if is_directory_host(page.get("url", "")):
        return 0
    full = normalize_key(name)
    core = normalize_key(core_company_name(name))
    title = normalize_key(page.get("title", ""))
    text = normalize_key(page.get("text", "")[:30000])
    score = 0
    if full and full in title:
        score = 100
    elif core and core in title:
        score = 80
    elif full and full in text:
        score = 70
    elif core and core in text:
        score = 50
    if full and full in text and score < 100:
        score += 10
    if core and core in text and score < 100:
        score += 5
    # 企业官网常见信息：地址/电话/邮箱/备案；命中年份、品牌词也略加分。
    for pat, bonus in (
        (r"地址|地\s*址|Address", 5),
        (r"电话|电\s*话|Tel|Phone", 5),
        (r"备案|ICP备|蜀ICP", 5),
        (r"关于我们|公司简介|About\s*Us", 5),
        (r"版权所有|Copyright", 3),
    ):
        if re.search(pat, page.get("text", ""), re.I):
            score += bonus
    return min(score, 120)


def site_confidence(score):
    """按候选首页得分分级：高分自动采用，中/低交由人工复核。"""
    if score >= 80:
        return "高（自动发现）"
    if score >= 50:
        return "中（自动发现，需复核）"
    return "低（自动发现，需复核）"


def discover_site(name, provided, args, confidence=""):
    """确认官网。有官网列时以用户提供为准；否则自动发现并分级置信度。

    返回值含逐候选分数（`candidates`），供 --discover-only 生成复核表。
    """
    if provided:
        url = norm_url(provided.split(",")[0].strip())
        if is_directory_host(url):
            # 用户或复核表填的是黄页/工商/名录站，不是企业官网：
            # 不采用该站，改为按公司名自动重新发现真实官网。
            log(f"[site] 忽略目录/聚合站（{url}），改为自动发现：{name}")
            provided = ""
    if provided:
        url = norm_url(provided.split(",")[0].strip())
        page, r = fetch_site(url, args)
        conf = confidence or "用户提供"
        if page:
            return {"url": page["url"], "confidence": conf, "page": page,
                    "error": "", "candidates": []}
        return {"url": url, "confidence": conf, "page": None,
                "error": r.get("error", "页面不可访问"), "candidates": []}
    candidates = search_candidates(name, args)
    scored, errors = [], []
    for url in candidates:
        page, r = fetch_site(url, args)
        if not page:
            errors.append(f"{url}: {r.get('error', '不可访问')}")
            continue
        score = site_score(name, page)
        if score:
            scored.append({"url": page.get("url") or url, "score": score,
                           "title": clean_inline(page.get("title", "")), "page": page})
    scored.sort(key=lambda x: -x["score"])
    if scored:
        best = scored[0]
        return {"url": best["url"], "confidence": site_confidence(best["score"]),
                "page": best["page"], "error": "", "score": best["score"],
                "candidates": [{k: v for k, v in c.items() if k != "page"} for c in scored]}
    return {"url": "", "confidence": "", "page": None,
            "error": "; ".join(errors[:3]), "candidates": []}


SITE_REVIEW_HEADERS = [
    "企业名称", "建议官网", "置信度", "分数", "候选官网", "决定", "自定义官网", "备注",
]
# 复核表「决定」列的识别词；留空＝未复核。
_SKIP_WORDS = {"跳过", "排除", "不采用", "否", "skip", "exclude", "no"}
_ADOPT_WORDS = {"采用", "确认", "是", "adopt", "yes", "y"}


def write_site_review(path, rows):
    """写官网候选复核表：决定列留空时，中/低置信度企业默认跳过。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    ws = wb.active
    ws.title = "官网复核"
    ws.append(SITE_REVIEW_HEADERS)
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    dv = DataValidation(type="list", formula1='"采用,跳过"', allow_blank=True,
                        showDropDown=False)
    ws.add_data_validation(dv)
    dv.add(f"F2:F{max(ws.max_row, 2)}")
    for col, width in zip("ABCDEFGH", (28, 34, 24, 8, 60, 10, 34, 24)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    ws2 = wb.create_sheet("使用说明")
    for line in (
        "1. 本表由 local_pipeline.py --discover-only 生成，列出各企业的自动发现候选与分数。",
        "2. 复核候选站点后，在「决定」列选择 采用 / 跳过；也可在「自定义官网」填正确网址（优先于「决定」）。",
        "3. 决定留空＝未复核：高置信度自动采用，中/低置信度默认跳过，不进入采集。",
        "4. 填好后运行：local_pipeline.py --excel <输入.xlsx> --site-decisions <本表.xlsx>",
        "5. 未出现在本表的企业按未复核处理，中/低置信度同样跳过。",
        "6. 候选分数规则见 references/pipeline.md 阶段 1：≥80 高，50-79 中，<50 低。",
    ):
        ws2.append([line])
    ws2.column_dimensions["A"].width = 110
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))


def load_site_decisions(path):
    """读回官网复核表的决定，返回 {企业名称: {建议官网, 置信度, 分数, 决定, 自定义官网}}。"""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb["官网复核"] if "官网复核" in wb.sheetnames else wb.active
        rows = ws.iter_rows(values_only=True)
        try:
            hdr = [str(c or "").strip() for c in next(rows)]
        except StopIteration:
            return {}

        def idx(name, fallback):
            return hdr.index(name) if name in hdr else fallback

        def cell(row, i):
            if i is None or i >= len(row) or row[i] is None:
                return ""
            return clean_inline(str(row[i]))

        i_name, i_site = idx("企业名称", 0), idx("建议官网", 1)
        i_conf, i_score = idx("置信度", 2), idx("分数", 3)
        i_dec, i_custom = idx("决定", 5), idx("自定义官网", 6)
        out = {}
        for row in rows:
            name = cell(row, i_name)
            if not name:
                continue
            out[name] = {
                "建议官网": cell(row, i_site),
                "置信度": cell(row, i_conf),
                "分数": cell(row, i_score),
                "决定": cell(row, i_dec),
                "自定义官网": cell(row, i_custom),
            }
        return out
    finally:
        wb.close()


def plan_site(item, decisions, args):
    """决定单家企业是否采集、用哪个官网。

    用户提供官网直接采用；复核表决定优先；未复核时只有高置信度自动采用，
    中/低置信度按“默认跳过”处理，不进入采集。
    """
    name = item["name"]
    provided = clean_inline(item.get("website", ""))
    if provided:
        return {"action": "crawl", "source": "用户提供",
                "discovery": discover_site(name, provided, args)}
    dec = decisions.get(name) or {}
    choice = clean_inline(str(dec.get("决定", ""))).lower()
    custom = clean_inline(str(dec.get("自定义官网", "")))
    if custom:
        # 显式填写自定义官网是最强人工信号，优先于「决定」列。
        return {"action": "crawl", "source": "人工确认",
                "discovery": discover_site(name, custom, args, confidence="人工确认")}
    if choice in _SKIP_WORDS:
        return {"action": "skip", "reason": "复核决定：跳过"}
    if choice in _ADOPT_WORDS:
        url = clean_inline(str(dec.get("建议官网", "")))
        if not url:
            return {"action": "skip", "reason": "复核决定采用，但表内没有建议官网"}
        return {"action": "crawl", "source": "人工确认",
                "discovery": discover_site(name, url, args, confidence="人工确认")}
    conf = clean_inline(str(dec.get("置信度", "")))
    if conf:
        # 复核表里已有结论：未填决定时只放行高置信度。
        if conf.startswith("高"):
            url = clean_inline(str(dec.get("建议官网", "")))
            if url:
                return {"action": "crawl", "source": "自动采用",
                        "discovery": discover_site(name, url, args, confidence=conf)}
        return {"action": "skip", "reason": f"未复核，自动发现置信度：{conf}"}
    discovered = discover_site(name, "", args)
    if discovered.get("url") and str(discovered.get("confidence", "")).startswith("高"):
        return {"action": "crawl", "source": "自动采用", "discovery": discovered}
    if not discovered.get("url") and item.get("resource_dir"):
        # 没有官网候选可复核时仍保留“仅用户资料成档”路径。
        return {"action": "crawl", "source": "仅用户资料", "discovery": discovered}
    return {"action": "skip",
            "reason": f"未复核，自动发现置信度：{discovered.get('confidence') or '未发现官网'}"}


def run_discover_only(companies, args, out):
    """两阶段流程的第一阶段：只做官网发现，产出候选复核表后退出。"""
    rows = []
    stats = collections.Counter()
    for i, item in enumerate(companies, 1):
        name = item["name"]
        provided = clean_inline(item.get("website", ""))
        d = discover_site(name, provided, args)
        conf = d.get("confidence", "")
        stats["用户提供" if provided else (conf or "未发现官网")] += 1
        cands = d.get("candidates") or []
        cand_text = "\n".join(f"{c.get('url', '')}（{c.get('score', '')}）" for c in cands[:8])
        if provided:
            note = "用户提供官网，无需复核"
        elif d.get("url"):
            note = "复核后填「决定」列；留空则中/低置信度默认跳过"
        else:
            note = (d.get("error", "") or "未发现候选官网")[:200]
        rows.append([name, d.get("url", ""), conf, d.get("score", ""),
                     cand_text, "", "", note])
        log(f"[{i}/{len(companies)}] {name} → {d.get('url') or '未发现'}"
            f"（{conf or '无'}）")
    target = out / "官网候选复核表.xlsx"
    if target.exists():
        target = out / f"官网候选复核表.{_dt.datetime.now().strftime('%Y%m%d-%H%M%S')}.xlsx"
    write_site_review(target, rows)
    log("\n官网发现统计：" + "；".join(f"{k} {v}" for k, v in stats.most_common()))
    log(f"复核表：{target}")
    log("请复核候选站点并在「决定」列填写 采用/跳过（留空＝未复核，中/低置信度将跳过），"
        "然后运行：")
    log(f'  local_pipeline.py --excel "<输入.xlsx>" --out "<输出目录>" '
        f'--site-decisions "{target}"')
    return 0


def sentence_list(text):
    text = clean_inline(text)
    parts = re.split(r"(?<=[。！？!?；;])\s*|。|！|？|;|；", text)
    out = []
    for p in parts:
        p = clean_inline(p)
        if 12 <= len(p) <= 180 and not NOISE.search(p):
            out.append(p.rstrip("，,；; ") + "。")
    return out


def extract_products(pages):
    candidates = []
    for page in pages:
        if page.get("cat") not in ("product", ""):
            continue
        candidates += page.get("headings", [])
        candidates += page.get("lists", [])
        if page.get("cat") == "product":
            candidates += [x.get("text", "") for x in page.get("links", [])]
    bad = re.compile(r"首页|关于|联系|导航|更多|查看|登录|注册|搜索|服务热线|在线留言|返回顶部|网站地图|人才招聘|新闻")
    columns = {"产品中心","行业应用","服务支持","投资者关系","关于我们","联系我们","新闻中心","人才招聘","企业文化","发展历程","资质荣誉","客户案例","解决方案","下载中心","在线留言","网站地图","供应商平台","业务咨询","品牌中心","应用领域","公司简介","企业简介","荣誉资质","客户应用","行业知识","全部","产品展示","新闻资讯","联系方式","招贤纳士","合作伙伴","在线客服","企业风采"}
    out, seen = [], set()
    for s in candidates:
        s = clean_inline(s).strip("·-—| ")
        s = re.sub(r"[\ue000-\uf8ff]", "", s).strip()
        if not 2 <= len(s) <= 50 or bad.search(s) or s in columns or NOISE.search(s):
            continue
        if re.search(r"电话|邮箱|地址|网址|@|\d{7,}", s):
            continue
        key = normalize_key(s)
        if key and key not in seen:
            seen.add(key)
            out.append(s)
    return out


def product_series(name):
    """取产品型号的"系列名"前缀，用于匹配首页省略完整型号的详情文本。"""
    m = re.search(r"^(.+?系列)", clean_inline(name))
    return m.group(1) if m else ""


def extract_product_details(pages, products):
    """从产品列表卡片链接文本提取"产品名 + 简介"里的简介部分。

    产品列表页的卡片链接通常形如"型号 完整简介 查看详情"，首页则可能只保留
    "系列名 + 简介"。返回 {产品名: 中文详情}，同一产品保留最长文本。
    """
    products = [clean_inline(p) for p in (products or []) if clean_inline(p)]
    details = {}
    exact = [(p, normalize_key(p)) for p in products]
    series = collections.defaultdict(list)
    for p in products:
        sp = product_series(p)
        if sp:
            series[normalize_key(sp)].append(p)

    def remember(product, detail):
        detail = clean_inline(detail).strip("：:，,-—·| ")
        if len(detail) < 12 or NOISE.search(detail):
            return
        if len(detail) > len(details.get(product, "")):
            details[product] = detail

    for page in pages:
        if page.get("cat") not in ("home", "product"):
            continue
        for link in page.get("links", []):
            text = DETAIL_TAIL.sub("", clean_inline(link.get("text", ""))).strip()
            if len(text) < 16:
                continue
            nkey = normalize_key(text)
            matched = ""
            for product, pkey in exact:
                if pkey and nkey.startswith(pkey) and len(pkey) > len(normalize_key(matched)):
                    matched = product
            if matched:
                remember(matched, text[len(matched):] if text.startswith(matched) else text)
                continue
            hits = []
            for spkey, plist in series.items():
                if len(spkey) >= 4 and nkey.startswith(spkey):
                    hits += plist
            if len(set(hits)) == 1 and hits[0] not in details:
                remember(hits[0], text)
    return details


def make_intro(pages, products):
    by_cat = collections.defaultdict(list)
    for p in pages:
        by_cat[p.get("cat", "home")].append(p.get("text", ""))
    overview = sentence_list(" ".join(by_cat.get("about", []) + by_cat.get("home", [])))
    business = sentence_list(" ".join(by_cat.get("product", [])))
    all_sent = []
    for p in pages:
        all_sent += sentence_list(p.get("text", ""))
    used = set()
    paras = []
    for group in (overview[:2], business[:2]):
        picked = []
        for s in group:
            k = normalize_key(s)
            if k not in used:
                used.add(k)
                picked.append(s)
        if picked:
            paras.append(" ".join(picked))
    if not business and products:
        paras.append("官网列出的主要产品包括：" + "、".join(products[:10]) + "。")
    for s in all_sent:
        k = normalize_key(s)
        if k not in used and len(paras) < 3:
            used.add(k)
            paras.append(s)
    return paras[:3]


def main_business(products, pages):
    for p in pages:
        for s in sentence_list(p.get("text", "")):
            if any(w in s for w in ("主营", "主要从事", "主要生产", "主要产品", "经营范围")):
                return s[:100]
    return "、".join(products[:3]) if products else ""


def extract_address(pages):
    for p in pages:
        for s in sentence_list(p.get("text", "")):
            if "地址" in s or "坐落" in s or "位于" in s:
                return s[:120]
    return ""


def guess_industry(text):
    for word in ("装备制造", "机械", "食品", "化工", "电子", "农业", "生物医药", "建材", "纺织", "新能源"):
        if word in text:
            return word
    return ""


def download_images(name, pages, home, args):
    buckets = {k: [] for k in IMAGE_DIRS}
    seen_hash = {}
    total = 0
    for page in pages:
        if total >= args.max_images:
            break
        page_cat = page.get("cat", "")
        # 目录/聚合站（黄页、工商信息、名录、B2B）整站素材与目标企业无关，
        # 在下载前直接跳过，避免广告图、二维码、无关缩略图进入交付。
        if is_directory_host(page.get("url", "")):
            continue
        for img in page.get("images", []):
            if total >= args.max_images:
                break
            cat = image_category(img, page_cat)
            if not cat:
                continue
            src = img.get("src", "")
            if not src.startswith(("http://", "https://")):
                continue
            if is_directory_host(src):
                continue
            r = fetch(src, args, binary=True, referer=page.get("url", ""))
            body = r.get("body") or b""
            if r.get("error") or not body or len(body) < 512 or len(body) > args.max_image_bytes:
                continue
            ct = ""
            try:
                ct = (r.get("headers").get_content_type() or "").lower()
            except Exception:
                pass
            if not ct.startswith("image/") and not re.search(r"\.(?:jpe?g|png|gif|webp|bmp|svg)(?:\?|$)", src, re.I):
                continue
            # 去重：同一内容只归档一次；产品图例外——同内容不同 alt 允许再存一份，
            # 以支持官网用同一张照片对应"美标/欧标"这类同图不同名的产品。
            # 任何情况下都禁止同一内容跨类别重复，避免工厂图/证书图互相串类。
            digest = hashlib.md5(body).hexdigest()
            alt_key = normalize_key(img.get("alt", ""))
            seen = seen_hash.setdefault(digest, {"cats": set(), "alts": set()})
            if cat != "product":
                if seen["cats"]:
                    continue
                seen["cats"].add(cat)
            else:
                if seen["cats"] - {"product"} or alt_key in seen["alts"]:
                    continue
                seen["cats"].add("product")
                seen["alts"].add(alt_key)
            ext = mimetypes.guess_extension(ct.split(";")[0]) or Path(urllib.parse.urlparse(src).path).suffix.lower()
            if ext in (".jpe",):
                ext = ".jpg"
            if ext not in (".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp", ".svg"):
                ext = ".jpg"
            stem = safe_filename(Path(urllib.parse.urlparse(src).path).stem, "image")[:40]
            folder = home / IMAGE_DIRS[cat]
            folder.mkdir(parents=True, exist_ok=True)
            filename = f"{len(buckets[cat]) + 1:02d}_{stem}{ext}"
            (folder / filename).write_bytes(body)
            buckets[cat].append({
                "file": filename, "alt": clean_inline(img.get("alt", "")),
                "url": src, "from": page.get("url", ""),
                "wh": "x".join(x for x in (img.get("width", ""), img.get("height", "")) if x),
            })
            total += 1
    return buckets


# ---- 用户资料摄入 --------------------------------------------------------
# 企业可能以任意文件夹结构、任意格式（图片/Word/Excel/PDF/PPT）提交资料。
# 本模块递归扫描资料根目录，按“扩展名 + 文件名/路径关键词”把素材归类，
# 图片直接归入四类，文档类抽取文本作为企业事实；原始文件另存备份。

RESOURCE_IMAGE_EXTS = (".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp", ".svg", ".tif", ".tiff")
RESOURCE_TEXT_EXTS = (".doc", ".docx", ".xls", ".xlsx", ".xlsm", ".pdf", ".ppt", ".pptx", ".txt", ".md", ".csv")
# 表格类只抽文本、不当作图片素材拷进图片文件夹（产品清单 xlsx 不应进入 2.企业产品图）。
RESOURCE_TABLE_EXTS = (".xls", ".xlsx", ".xlsm", ".csv")

# 关键词 -> 归类。优先匹配更具体的类别；大小写不敏感。
RESOURCE_CAT_KEYWORDS = [
    ("logo", ("logo", "标志", "徽标", "商标", "标识", "brand")),
    ("cert", ("证书", "资质", "荣誉", "认证", "certificate", "cert", "honor", "award",
              "iso", "专利", "许可", "检测报告", "营业执照")),
    ("factory", ("工厂", "车间", "厂房", "厂区", "生产", "设备", "基地", "生产线",
                 "factory", "workshop", "plant", "equipment")),
    ("product", ("产品", "商品", "产品图", "样品", "型号", "product", "goods", "pro_")),
    ("about", ("简介", "介绍", "关于", "公司", "profile", "about", "company", "宣传")),
]


def resource_category(name):
    """按文件名/相对路径关键词返回归类：logo/cert/factory/product/about/''。"""
    s = clean_inline(name).lower()
    for cat, words in RESOURCE_CAT_KEYWORDS:
        if any(w in s for w in words):
            return cat
    return ""


def _docx_text(path):
    try:
        from docx import Document
        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
                if cells:
                    parts.append(" ".join(cells))
        return "\n".join(parts)
    except Exception:
        return ""


def _xlsx_text(path):
    try:
        from openpyxl import load_workbook
        wb = load_workbook(str(path), read_only=True, data_only=True)
        parts = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                cells = [clean_inline(c) for c in row if c not in (None, "")]
                if cells:
                    # 用制表符保留列边界，供产品清单按“产品名/型号/说明”分列解析。
                    parts.append("\t".join(cells))
        return "\n".join(parts)
    except Exception:
        return ""


def _pptx_text(path):
    """PPTX 文本走 zipfile + XML，避免额外依赖。"""
    try:
        import zipfile
        from xml.etree import ElementTree as ET
        ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
        parts = []
        with zipfile.ZipFile(str(path)) as z:
            slides = sorted(n for n in z.namelist()
                            if re.match(r"ppt/slides/slide\d+\.xml$", n))
            for name in slides:
                try:
                    root = ET.fromstring(z.read(name))
                except Exception:
                    continue
                texts = [t.text for t in root.iter(ns + "t") if t.text and t.text.strip()]
                if texts:
                    parts.append(" ".join(texts))
        return "\n".join(parts)
    except Exception:
        return ""


def _pdf_text(path):
    """PDF 抽取尽力而为：环境有 pypdf/PyPDF2 才抽，否则只归档。"""
    for mod in ("pypdf", "PyPDF2"):
        try:
            m = __import__(mod)
            reader = m.PdfReader(str(path))
            parts = []
            for page in reader.pages[:20]:
                try:
                    t = page.extract_text() or ""
                except Exception:
                    t = ""
                if t.strip():
                    parts.append(t)
            return "\n".join(parts)
        except Exception:
            continue
    return ""


def resource_text(path):
    """按扩展名抽取文档文本；不支持或失败返回空串。"""
    ext = Path(path).suffix.lower()
    if ext == ".docx":
        return _docx_text(path)
    if ext in (".xlsx", ".xlsm"):
        return _xlsx_text(path)
    if ext == ".pptx":
        return _pptx_text(path)
    if ext == ".pdf":
        return _pdf_text(path)
    if ext in (".txt", ".md", ".csv"):
        for enc in ("utf-8", "gbk", "utf-16"):
            try:
                return Path(path).read_text(encoding=enc)
            except Exception:
                continue
    return ""


def resource_match_tokens(name):
    """企业名的匹配 token：全称、去地域/去后缀核心名、以及核心名去掉通用词后的品牌词。"""
    full = normalize_key(name)
    core = normalize_key(core_company_name(name))
    toks = set()
    for t in (full, core):
        if len(t) >= 3:
            toks.add(t)
    generic = ("新能源科技", "科技", "机械制造", "制造", "实业", "贸易", "电子商务",
               "电子", "材料", "装备", "工程", "食品", "农业", "生物", "医药", "环保",
               "设备", "建设", "发展", "管理", "服务", "文化", "旅游", "物流")
    for t in (core, full):
        for g in generic:
            if t.endswith(g) and len(t) - len(g) >= 2:
                toks.add(t[:-len(g)])
            if g in t and len(t) - len(g) >= 2:
                toks.add(t.replace(g, ""))
    # 前缀品牌词：资料目录常只写企业名开头（如“川澜”“天翔”），补 core 的前缀子串。
    # 只取 2-4 字前缀，且不落入通用词，降低误匹配。
    for n in (2, 3, 4):
        if len(core) > n:
            toks.add(core[:n])
    toks -= set(generic)
    return {t for t in toks if len(t) >= 2}


def match_resource_dir(companies, root):
    """把资料根目录下的子文件夹匹配到企业名。

    规则：子文件夹名与企业全称/核心名/品牌词互相包含即算命中；子文件夹名含
    “资料/文件/材料”等无意义词时先剥离。多家企业命中同一目录时按最长 token 归属。
    返回 {企业名: 企业资料目录}，未命中的企业不出现在结果里。
    """
    result = {}
    if not root or not Path(root).is_dir():
        return result
    root = Path(root)
    subdirs = [d for d in sorted(root.iterdir()) if d.is_dir()]
    stop = ("资料", "文件", "素材", "图片", "照片", "文档", "企业", "公司")
    keys = {c: resource_match_tokens(c) for c in companies}
    for d in subdirs:
        dkey = normalize_key(d.name)
        if not dkey:
            continue
        for w in stop:
            dkey = dkey.replace(w, "")
        if not dkey:
            continue
        best, best_len = "", 0
        for c, toks in keys.items():
            for t in toks:
                if t and (t in dkey or dkey in t) and len(t) > best_len:
                    best, best_len = c, len(t)
        if best:
            result[best] = d
    return result


def ingest_resources(company, resource_dir, raw_home, backup_home, args):
    """扫描一家企业的资料目录，返回 (updates, notes)。

    - 图片：按类别直接拷入 raw_home/<类别目录>，文件名以 user_ 前缀区分来源；
      render 阶段会把它连同官网图一起搬进交付目录的五文件夹。
    - 文档：抽文本，归入简介/产品/资质事实；原始文件复制到 backup_home 备份。
      表格类（xls/xlsx/xlsm/csv）只抽文本，不拷进图片文件夹。
    - 未知格式：原件照样备份并记入未归类清单，不静默丢弃。
    """
    updates = {"logo": [], "product": [], "cert": [], "factory": []}
    notes = {"企业": company, "文档": [], "表格": [], "未归类": [], "抽取失败": []}
    if not resource_dir or not Path(resource_dir).is_dir():
        return updates, notes
    resource_dir = Path(resource_dir)
    raw_home = Path(raw_home)
    if str(backup_home):
        backup_home = Path(backup_home)
        backup_home.mkdir(parents=True, exist_ok=True)
    doc_texts = []
    for fp in sorted(resource_dir.rglob("*")):
        if not fp.is_file():
            continue
        rel = fp.relative_to(resource_dir)
        rel_key = str(rel)  # 用相对路径参与关键词判断（目录名也可能含类别词）
        ext = fp.suffix.lower()
        # 原始文件一律备份（在交付父文件夹之外保留用户交来的全部资料）
        try:
            dest = backup_home / rel
            dest.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(fp, dest)
        except Exception as exc:
            notes["抽取失败"].append(f"备份失败 {rel}: {exc}")
        if ext in RESOURCE_IMAGE_EXTS:
            cat = resource_category(rel_key)
            if cat in updates:
                folder = raw_home / IMAGE_DIRS[cat]
                folder.mkdir(parents=True, exist_ok=True)
                filename = f"user_{safe_filename(fp.stem)}{ext}"
                try:
                    shutil.copy2(fp, folder / filename)
                except Exception as exc:
                    notes["抽取失败"].append(f"图片拷贝失败 {rel}: {exc}")
                    continue
                updates[cat].append({
                    "file": filename, "alt": clean_inline(fp.stem),
                    "url": "", "from": str(rel), "wh": "", "source": "用户资料",
                })
            else:
                notes["未归类"].append(f"图片未归类：{rel}")
        elif ext in RESOURCE_TEXT_EXTS:
            cat = resource_category(rel_key)
            text = resource_text(fp)
            if text.strip():
                doc_texts.append({"rel": str(rel), "cat": cat, "text": text})
            elif ext == ".pdf":
                notes["抽取失败"].append(f"PDF 未抽到文本（已备份）：{rel}")
            # 文档/PDF/PPT 等只抽取文本或事实，原件归入交付之外的备份目录；
            # 交付五文件夹内只放图片素材，避免非图片文件成为未登记孤儿。
            if ext in RESOURCE_TABLE_EXTS:
                notes["表格"].append({"文件": str(rel), "归类": cat or "待定", "字符数": len(text)})
        else:
            notes["未归类"].append(f"未知格式：{rel}")
    for d in doc_texts:
        notes["文档"].append({"文件": d["rel"], "归类": d["cat"] or "待定", "字符数": len(d["text"])})
    updates["_doc_texts"] = doc_texts
    return updates, notes


def resource_notes_summary(notes):
    """把摄入备注压成一句人类可读的摘要，写入 raw 与汇总表。"""
    if not notes:
        return ""
    docs = len(notes.get("文档") or [])
    tables = len(notes.get("表格") or [])
    unclassified = len(notes.get("未归类") or [])
    failed = len(notes.get("抽取失败") or [])
    parts = []
    if docs:
        parts.append(f"抽取文档 {docs} 份")
    if tables:
        parts.append(f"抽到表格 {tables} 份")
    if unclassified:
        parts.append(f"未归类 {unclassified} 项")
    if failed:
        parts.append(f"抽取/备份失败 {failed} 项")
    return "用户资料：" + ("、".join(parts) if parts else "已摄入")


def _docs_to_facts(doc_texts):
    """从用户文档文本中提取简介段落、产品名与产品详情。"""
    intro, products, details = [], [], {}
    seen_intro, seen_product = set(), set()
    for d in doc_texts:
        raw = d.get("text", "")
        table_like = "\t" in raw or d.get("table") is True
        text = clean_lines(raw) if table_like else clean_inline(raw)
        cat = d.get("cat", "")
        # 简介：优先 about/简介类文档，或含明显公司介绍句的文本
        for s in sentence_list(text):
            if len(intro) >= 3:
                break
            if any(w in s for w in ("公司", "企业", "成立", "位于", "主营", "是一家", "简介", "工厂")):
                k = normalize_key(s)
                if k and k not in seen_intro:
                    seen_intro.add(k)
                    intro.append(s)
        # 产品：只从产品类文档或含明确产品清单的结构行里提；一行一个产品，
        # 取该行第一个短字段作为产品名，整行作为产品详情，避免把简介句/导航词当产品。
        looks_like_table = ("产品" in text[:80] or "型号" in text[:80] or "系列" in text[:80])
        if cat == "product" or looks_like_table:
            for line in text.split("\n"):
                line = line.strip()
                if not line:
                    continue
                has_col = "\t" in line
                # 整行纯叙述句（含句号且很长、且不是分列表格行）不是产品行，跳过。
                if not has_col and len(line) > 40 and ("。" in line or "，" in line):
                    continue
                cells = [x.strip() for x in (line.split("\t") if has_col
                                             else re.split(r"[，、;；|]+|\s{2,}", line))]
                cells = [x for x in cells if x]
                if not cells:
                    continue
                name = cells[0]
                # 表头行/说明行过滤
                if name in ("产品名称", "名称", "产品", "型号", "产品名"):
                    continue
                if not (2 <= len(name) <= 20):
                    continue
                if any(w in name for w in ("公司", "企业", "成立于", "位于", "在线", "主营")):
                    continue
                k = normalize_key(name)
                if not k or k in seen_product:
                    continue
                seen_product.add(k)
                products.append(name)
                # 产品名/型号行本身不是详情；只有行内确有额外描述时才记录。
                if normalize_key(line) != normalize_key(name):
                    details.setdefault(name, line[:200])
    return intro, products, details


def finalize_resource_only(archive, res_updates, args):
    """官网不可用、仅凭用户资料成档时，填充事实与图片。"""
    doc_texts = res_updates.get("_doc_texts") or []
    intro, products, details = _docs_to_facts(doc_texts)
    for k in IMAGE_DIRS:
        archive[k] = list(res_updates.get(k) or [])
    if products:
        archive["products"] = products[:args.max_products] if args.max_products else products
    if details:
        archive["product_details"] = details
    if not intro:
        # 无 about 文档时，用最长的一段资料文本兜底，避免简介全空。
        blob = max((clean_inline(d.get("text", "")) for d in doc_texts), key=len, default="")
        for s in sentence_list(blob):
            if len(s) >= 20:
                intro.append(s)
            if len(intro) >= 3:
                break
    archive["intro_paragraphs"] = intro[:3]
    archive["核心产品"] = "、".join(archive.get("products", [])[:8])
    archive["主要业务"] = archive.get("主要业务") or "、".join(archive.get("products", [])[:6])
    return archive


def merge_user_resources(archive, res_updates, buckets):
    """用户资料优先、官网补充：图片前置用户素材，文档事实覆盖官网推断。"""
    doc_texts = res_updates.get("_doc_texts") or []
    user_intro, user_products, user_details = _docs_to_facts(doc_texts)
    # 图片：用户素材在前，官网素材在后；同一文件名不重复。
    for k in IMAGE_DIRS:
        user_imgs = list(res_updates.get(k) or [])
        site_imgs = list(buckets.get(k) or [])
        seen = {it.get("file") for it in user_imgs}
        merged = user_imgs + [it for it in site_imgs if it.get("file") not in seen]
        archive[k] = merged
    # 文本事实：用户资料优先
    if user_intro:
        archive["intro_paragraphs"] = (user_intro + [p for p in archive.get("intro_paragraphs", [])
                                                       if normalize_key(p) not in {normalize_key(x) for x in user_intro}])[:3]
    if user_products:
        # 用户资料优先：官网产品仅作补充，且仅在官网可信时并入，避免中/低置信度误匹配站污染清单。
        confidence = archive.get("置信度", "") or ""
        site_ok = not (confidence.startswith("中") or confidence.startswith("低") or "需复核" in confidence)
        seen = {normalize_key(p) for p in user_products}
        extra = [p for p in archive.get("products", [])
                 if site_ok and normalize_key(p) not in seen] if site_ok else []
        archive["products"] = user_products + extra
    if user_details:
        d = dict(archive.get("product_details") or {})
        d.update(user_details)
        archive["product_details"] = d
    archive["核心产品"] = "、".join(archive.get("products", [])[:8])
    return archive


# ---- 用户资料摄入结束 ----------------------------------------------------


def crawl_company(item, args, out_root):
    name = item["name"]
    archive = {
        "名称": name, "官网": "", "置信度": "", "处理地": "", "产业": "", "规模": "",
        "主要业务": "", "地址": "", "核心产品": "", "logo": [], "product": [],
        "cert": [], "factory": [], "pages": [], "products": [], "errors": [],
        "about_text": "", "product_text": "", "home_text": "", "intro_paragraphs": [],
        "product_details": {}, "status": "partial", "状态原因": "",
        "资料备注": "", "资料来源": "",
    }
    # 先摄入用户资料：官网不可用时资料仍可产出，且资料优先级高于官网。
    backup_root = Path(getattr(args, "backup_dir", "") or "")
    resource_dir = (item.get("resource_dir") or "")
    res_updates, res_notes = {"logo": [], "product": [], "cert": [], "factory": []}, {}
    if resource_dir:
        raw_home = out_root / safe_filename(name)
        bhome = (backup_root / safe_filename(name)) if str(backup_root) else Path("")
        res_updates, res_notes = ingest_resources(name, resource_dir, raw_home, bhome, args)
        archive["资料来源"] = str(resource_dir)
        archive["资料备注"] = resource_notes_summary(res_notes)
    # 两阶段流程：发现/复核结果由 main 预置，避免重复搜索。
    discovered = item.get("_site_discovery") or discover_site(name, item.get("website", ""), args)
    if discovered.get("url") and is_directory_host(discovered["url"]):
        # 兜底：任何来源（用户提供/复核表/自动发现）落在目录聚合站上都不采集。
        log(f"[site] 确认官网为目录/聚合站，放弃采集：{discovered['url']}")
        discovered = {"url": "", "confidence": "",
                      "page": None, "error": f"目录/聚合站不作为官网：{discovered['url']}",
                      "candidates": []}
    archive["官网"] = discovered.get("url", "")
    archive["置信度"] = discovered.get("confidence", "")
    if discovered.get("error"):
        archive["errors"].append(discovered["error"])
    has_resources = bool(res_updates.get("_doc_texts") or res_updates.get("logo")
                         or res_updates.get("product") or res_updates.get("cert")
                         or res_updates.get("factory"))
    if not discovered.get("url"):
        if has_resources:
            # 官网没找到，但用户资料可用：以资料成档。
            archive["status"] = "resource_only"
            archive["状态原因"] = "未确认官网，基于用户提供资料成档"
            return finalize_resource_only(archive, res_updates, args)
        archive["status"] = "no_website"
        archive["状态原因"] = "本地发现与校验未确认官网"
        return archive
    home_page = discovered.get("page")
    if not home_page:
        home_page, r = fetch_page(archive["官网"], args)
        if not home_page:
            if has_resources:
                archive["status"] = "resource_only"
                archive["状态原因"] = "官网不可访问，基于用户提供资料成档"
                return finalize_resource_only(archive, res_updates, args)
            archive["status"] = "no_website"
            archive["状态原因"] = "官网页面不可访问：" + r.get("error", "")
            return archive
    pages = [{"cat": "home", **home_page}]
    archive["home_text"] = home_page.get("text", "")

    def page_kind(url):
        """detail = 单条产品详情页；index = 分类/列表页（一次带出多个产品）。"""
        path = urllib.parse.urlparse(url).path.lower()
        return "detail" if re.search(r"\.(?:html?|php|aspx?|jsp)$", path) else "index"

    def site_links(page):
        out = []
        for link in page.get("links", []):
            url, text = link.get("href", ""), link.get("text", "")
            if not url or not same_site(url, page.get("url", "")):
                continue
            cat = url_category(url, text)
            if cat:
                out.append((url, cat, text))
        return out

    seen = {home_page.get("url", "")}
    home_links = site_links(home_page)
    # 产品分类/列表页优先于单条详情页：列表页能一次带出多个产品名、详情与产品图。
    ordered = ([x for x in home_links if x[1] == "product" and page_kind(x[0]) == "index"]
               + [x for x in home_links if x[1] != "product"]
               + [x for x in home_links if x[1] == "product" and page_kind(x[0]) == "detail"])
    selected = []
    for url, cat, text in ordered:
        if url in seen:
            continue
        seen.add(url)
        selected.append((url, cat))
        if len(selected) >= args.max_pages - 1:
            break
    for url, cat in selected:
        page, r = fetch_page(url, args)
        if not page:
            archive["errors"].append(f"{url}: {r.get('error', '不可访问')}")
            continue
        page["cat"] = cat
        pages.append(page)
        if cat == "about":
            archive["about_text"] = page.get("text", "")
        elif cat == "product":
            archive["product_text"] = page.get("text", "")
    # 再下钻一层：从已抓产品页补齐叶子分类与产品详情页，分类页优先。
    extra_index, extra_detail = [], []
    for page in pages:
        if page.get("cat") != "product":
            continue
        for url, cat, text in site_links(page):
            if url in seen:
                continue
            (extra_index if cat == "product" and page_kind(url) == "index" else extra_detail).append((url, cat))
    for url, cat in extra_index + extra_detail:
        if len(pages) >= args.max_pages:
            break
        if url in seen:
            continue
        seen.add(url)
        page, r = fetch_page(url, args)
        if not page:
            archive["errors"].append(f"{url}: {r.get('error', '不可访问')}")
            continue
        page["cat"] = cat
        pages.append(page)
    archive["pages"] = [{"cat": p.get("cat", ""), "url": p.get("url", ""), "title": p.get("title", "")} for p in pages]
    all_products = extract_products(pages)
    details = extract_product_details(pages, all_products)
    home = out_root / safe_filename(name)
    for folder in IMAGE_DIRS.values():
        (home / folder).mkdir(parents=True, exist_ok=True)
    buckets = download_images(name, pages, home, args)
    product_image_keys = [normalize_key(it.get("alt", "")) for it in (buckets.get("product") or [])]

    def has_product_evidence(product):
        """只有产品详情或产品图 alt 明确佐证的条目才算产品，过滤导航/分类词。"""
        if product in details:
            return True
        pkey = normalize_key(product)
        if not pkey:
            return False
        spkey = normalize_key(product_series(product))
        for key in product_image_keys:
            if not key:
                continue
            if key == pkey:
                return True
            # 系列名允许与图片 alt 互为前缀，避免把"储能系统"这类分类词当产品。
            if spkey and len(spkey) >= 4 and (key.startswith(spkey) or spkey.startswith(key)):
                return True
        return False

    archive["products"] = [p for p in all_products if has_product_evidence(p)]
    archive["product_details"] = {k: v for k, v in details.items() if k in archive["products"]}
    archive["intro_paragraphs"] = make_intro(pages, archive["products"])
    archive["主要业务"] = main_business(archive["products"], pages)
    archive["产业"] = guess_industry(" ".join(p.get("text", "") for p in pages))
    archive["地址"] = extract_address(pages)
    archive["核心产品"] = "、".join(archive["products"][:8])
    # 用户资料优先、官网补充：图片前置用户素材，文档事实覆盖官网推断。
    merge_user_resources(archive, res_updates, buckets)
    total_images = sum(len(archive.get(k) or []) for k in IMAGE_DIRS)
    if total_images == 0:
        archive["status"] = "empty_images"
        archive["状态原因"] = "官网与用户资料均未提供可归档图片"
    elif archive["errors"] or len(archive["intro_paragraphs"]) < 3:
        archive["status"] = "partial"
        archive["状态原因"] = "抓取存在错误或中文简介不足三段"
    else:
        archive["status"] = "ok"
        archive["状态原因"] = ""
    return archive


def write_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def make_docx(path, archive, en_entry):
    from docx import Document
    doc = Document()
    title = archive["名称"]
    if archive.get("主要业务"):
        title += f"（{archive['主要业务']}）"
    doc.add_paragraph(title)
    paras = archive.get("intro_paragraphs") or [archive.get("home_text", "")[:500] or "官网公开信息不足，待补充。"]
    for p in paras[:3]:
        doc.add_paragraph(p)
    if en_entry:
        for p in (en_entry.get("英文简介") or [])[:3]:
            if isinstance(p, str) and p.strip():
                doc.add_paragraph(p.strip())
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def _product_match_keys(product):
    """产品名的可匹配键：整名/主体=强键(2)，型号 token=弱键(1)。

    返回 {key: weight}。强键用于完整产品族匹配，弱键只在精确相等时才够用，
    避免多个型号共享同一段后缀（如 10X100T 退化成 X100T）而一张图错挂多行。
    """
    text = clean_inline(product)
    keys = {}

    def add(key, weight):
        if len(key) >= 2:
            keys[key] = max(keys.get(key, 0), weight)

    for value in (text, product_series(text)):
        add(normalize_key(value), 2)
    # “储能柜 型号 ESS-100”这类名称，型号前的主体是有效的产品族匹配词。
    base = re.split(r"\s*(?:型号|model|spec(?:ification)?)\s*", text, flags=re.I)[0]
    add(normalize_key(base), 2)
    # 型号整体抽取（字母数字混合，含 10X100T）；裸数字片段不再当作键。
    for token in re.findall(r"[A-Za-z0-9][A-Za-z0-9._-]*[A-Za-z][A-Za-z0-9._-]*\d[A-Za-z0-9._-]*", text):
        add(normalize_key(token), 1)
    return keys


def _product_image_keys(item):
    """从图片 alt、文件名、来源路径和 URL 提取匹配文本。"""
    keys = set()
    texts = [
        item.get("alt", ""),
        Path(clean_inline(item.get("file", ""))).stem,
        item.get("from", ""),
        item.get("url", ""),
    ]
    for text in texts:
        text = clean_inline(text)
        if not text:
            continue
        pieces = [text] + re.split(r"[\\/|]+", text)
        for piece in pieces:
            piece = re.sub(r"^user[_-]+", "", piece, flags=re.I)
            piece = re.sub(r"\.[A-Za-z0-9]{1,5}$", "", piece)
            key = normalize_key(piece)
            if len(key) >= 2:
                keys.add(key)
            for token in re.findall(r"[\u4e00-\u9fff]{2,}|[A-Za-z0-9][A-Za-z0-9._-]*\d[A-Za-z0-9._-]*", piece):
                token_key = normalize_key(token)
                if len(token_key) >= 2:
                    keys.add(token_key)
    return keys


def _product_image_match_score(product_keys, candidate_keys):
    """打分：强键(整名/主体)精确=4、弱键(型号)精确=3、强键包含=2。

    只允许强键做包含匹配；弱型号键必须精确相等，避免共享后缀串图。
    """
    best = 0
    for pkey, weight in product_keys.items():
        for ckey in candidate_keys:
            if pkey == ckey:
                best = max(best, 3 + (1 if weight >= 2 else 0))
            elif weight >= 2 and len(pkey) >= 3 and len(ckey) >= 3 and (pkey in ckey or ckey in pkey):
                best = max(best, 2)
    return best


def product_image_links(archive):
    """按图片元数据匹配产品名，返回 {产品名: [本地相对路径]}。

    匹配优先级：alt/文件名/来源路径/URL 完全一致 > 产品主体或型号包含匹配。
    用户资料只有一张产品图且无法精确匹配时，作为未匹配产品行的主图兜底，
    保证产品清单仍有可核验的本地链接；多张图无法唯一匹配时不强行复用。
    """
    folder = IMAGE_DIRS["product"]
    candidates = []
    for item in archive.get("product") or []:
        file_name = clean_inline(item.get("file", ""))
        if not file_name:
            continue
        keys = _product_image_keys(item)
        if not keys:
            continue
        candidates.append({
            "rel": f"{folder}/{file_name}",
            "keys": keys,
            "source": clean_inline(item.get("source", "")),
        })
    result = {}
    products = list(archive.get("products") or [])
    product_keys = {product: _product_match_keys(product) for product in products}
    for product in products:
        scored = []
        for cand in candidates:
            score = _product_image_match_score(product_keys[product], cand["keys"])
            if score:
                scored.append((score, cand))
        if scored:
            best_score = max(x[0] for x in scored)
            matches = [x[1] for x in scored if x[0] == best_score]
        else:
            matches = []
        rels, seen_rel = [], set()
        for cand in matches:
            if cand["rel"] not in seen_rel:
                seen_rel.add(cand["rel"])
                rels.append(cand["rel"])
        result[product] = rels

    user_candidates = [c for c in candidates if c["source"] == "用户资料"]
    unmatched = [p for p in products if not result.get(p)]
    if len(user_candidates) == 1 and unmatched:
        rel = user_candidates[0]["rel"]
        for product in unmatched:
            result[product] = [rel]
    return result


def make_product_xlsx(path, archive, en_entry):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = "产品清单"
    ws.append(PRODUCT_COLUMNS)
    pm = (en_entry or {}).get("产品英名") or {}
    dm = (en_entry or {}).get("产品详情英") or {}
    details = archive.get("product_details") or {}
    brand = clean_inline((en_entry or {}).get("品牌", ""))
    products = archive.get("products", [])
    image_links = product_image_links(archive)
    detail_hit = 0
    for i, product in enumerate(products, 1):
        eng = clean_inline(pm.get(product, "")) if isinstance(pm, dict) else ""
        cell = f"{product} / {eng}" if eng else product
        zh_detail = clean_inline(details.get(product, ""))
        en_detail = clean_inline(dm.get(product, "")) if isinstance(dm, dict) else ""
        if zh_detail:
            detail_hit += 1
            detail_cell = f"{zh_detail}\n{en_detail}" if en_detail else zh_detail
        else:
            detail_cell = ""
        links = image_links.get(product) or []
        # 来源没有对应产品图时直接留空，不写占位说明。
        img_cell = "\n".join(links) if links else ""
        ws.append([archive.get("产业", ""), i, cell, brand, "", "", detail_cell, img_cell])
        if links:
            img_cell_obj = ws.cell(row=ws.max_row, column=8)
            img_cell_obj.hyperlink = links[0]
            img_cell_obj.font = Font(color="0563C1", underline="single")
    ws2 = wb.create_sheet("企业说明")
    ws2.append(["企业名称", "英文名称", "源文件夹", "产品数", "产品详情覆盖", "资料备注", "价格说明"])
    ws2.append([
        archive["名称"], (en_entry or {}).get("英文名", ""), safe_filename(archive["名称"]),
        len(products), f"{detail_hit}/{len(products)}", archive.get("状态原因", ""), "",
    ])
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))


def make_summary(path, archives, en_data):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = "官网汇总"
    ws.append(SUMMARY_HEADERS)
    for a in archives:
        en = en_data.get(a["名称"], {})
        ws.append([
            a["名称"], a.get("官网", ""), a.get("置信度", ""), "", a.get("产业", ""),
            a.get("规模", ""), a.get("主要业务", ""), a.get("地址", ""), a.get("核心产品", ""),
            sum(len(a.get(k) or []) for k in IMAGE_DIRS),
            len(a.get("logo") or []), len(a.get("product") or []), len(a.get("cert") or []),
            len(a.get("factory") or []), len(a.get("products") or []),
            len(a.get("intro_paragraphs") or []), en.get("英文名", ""), a.get("官网", ""), "是",
        ])
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))


def load_en(path, archives, args=None, cache_path=None):
    """有 --en 用人工确认稿；否则默认自动翻译生成中英双语草稿。"""
    if path and Path(path).is_file():
        data = json.loads(Path(path).read_text(encoding="utf-8"))
        if not isinstance(data, dict):
            raise SystemExit("英文文件必须是 JSON 对象")
        for entry in data.values():
            if isinstance(entry, dict):
                entry.setdefault("定稿", True)
                entry.setdefault("翻译失败", False)
        return data
    if args is None or getattr(args, "no_translate", False):
        return {a["名称"]: empty_english_entry(a) for a in archives}
    translator = Translator(args, cache_path)
    log(f"自动生成英文：{len(archives)} 家（MyMemory 中译英，结果标为待人工核校）")
    data = {}
    for i, a in enumerate(archives, 1):
        entry = auto_english_entry(a, translator)
        data[a["名称"]] = entry
        log(f"  英文[{i}/{len(archives)}] {a['名称']} · "
            f"{'失败' if entry['翻译失败'] else 'ok'} · "
            f"简介 {len(entry['英文简介'])} 段 · 产品英名 {len(entry['产品英名'])}")
    translator.save()
    return data


def render_stage(stage, archives, en_data):
    deliverable = stage / "deliverable"
    raw_dir = stage / "raw"
    raw_dir.mkdir(parents=True, exist_ok=True)
    for a in archives:
        write_json(raw_dir / (safe_filename(a["名称"]) + ".json"), a)
    for a in archives:
        home = deliverable / safe_filename(a["名称"])
        source_home = stage / "raw_home" / safe_filename(a["名称"])
        for key, folder in IMAGE_DIRS.items():
            target = home / folder
            target.mkdir(parents=True, exist_ok=True)
            source = source_home / folder
            if source.is_dir():
                for image in source.iterdir():
                    if image.is_file():
                        shutil.copy2(image, target / image.name)
        intro_dir = home / "5.企业介绍"
        intro_dir.mkdir(parents=True, exist_ok=True)
        for stale in intro_dir.glob(f"{safe_filename(a['名称'])}简介*短*.docx"):
            try:
                stale.unlink()
            except OSError:
                pass
        make_docx(intro_dir / f"{safe_filename(a['名称'])}.docx", a, en_data.get(a["名称"]))
        make_product_xlsx(home / "产品清单.xlsx", a, en_data.get(a["名称"]))
    write_json(stage / "en.json", en_data)
    make_summary(stage / "汇总.xlsx", archives, en_data)
    return deliverable, raw_dir, stage / "en.json", stage / "汇总.xlsx"


def build_visual_review(stage, deliverable, raw_dir):
    """生成拼版、缩略图核对表和视觉核对.json 结论载体。"""
    script = Path(__file__).with_name("visual_review.py")
    if not script.is_file():
        return None, "找不到 visual_review.py"
    json_path = stage / "视觉核对.json"
    cmd = [
        sys.executable, str(script), "--deliverable", str(deliverable), "--raw", str(raw_dir),
        "--out", str(stage / "review"), "--json", str(json_path),
    ]
    proc = subprocess.run(cmd, text=True, encoding="utf-8", errors="replace", capture_output=True)
    log_text = proc.stdout + ("\n" + proc.stderr if proc.stderr else "")
    (stage / "视觉核对.log").write_text(log_text, encoding="utf-8")
    if proc.returncode != 0:
        return None, log_text
    return json_path, log_text


def run_gates(stage, deliverable, raw_dir, en_path, summary, expected, visual=None,
              require_visual=True, allow_no_english=False, allow_builder_cdn=False):
    gate = Path(__file__).with_name("gates.py")
    if not gate.is_file():
        return None, "找不到 gates.py"
    cmd = [
        sys.executable, str(gate), "--deliverable", str(deliverable), "--raw", str(raw_dir),
        "--en", str(en_path), "--summary", str(summary), "--expected", str(expected),
        "--json", str(stage / "gates.json"),
    ]
    if visual:
        cmd += ["--visual", str(visual)]
    cmd += ["--require-visual" if require_visual else "--skip-visual-review"]
    if allow_no_english:
        cmd += ["--accept-no-english"]
    if allow_builder_cdn:
        cmd += ["--allow-builder-cdn"]
    proc = subprocess.run(cmd, text=True, encoding="utf-8", errors="replace", capture_output=True)
    (stage / "gates.log").write_text(proc.stdout + ("\n" + proc.stderr if proc.stderr else ""), encoding="utf-8")
    return proc.returncode == 0, proc.stdout + ("\n" + proc.stderr if proc.stderr else "")


def publish(stage, out, run_id, target=None):
    """把 stage 内容发布到目标目录，返回实际发布目录。

    目标里同名文件被占用（最常见的是上一轮 xlsx 还开在 Excel 中）时，
    copy 会抛 OSError；此时整批改发到 <out>/publish_<run_id>/，
    保证“产物已生成且门禁通过”不会因为一个被锁文件整条崩掉。
    """
    target = Path(target) if target else out
    target.mkdir(parents=True, exist_ok=True)

    def copy_all(dest):
        for item in stage.iterdir():
            dst = dest / item.name
            if item.is_dir():
                shutil.copytree(item, dst, dirs_exist_ok=True)
            else:
                shutil.copy2(item, dst)

    try:
        copy_all(target)
    except OSError as exc:
        fallback = out / f"publish_{run_id}"
        log(f"发布目录被占用（{exc}），改发到：{fallback}")
        fallback.mkdir(parents=True, exist_ok=True)
        copy_all(fallback)
        target = fallback
    write_json(target / "manifest.json", {
        "run_id": run_id, "published_at": _dt.datetime.now().isoformat(timespec="seconds"),
        "source_run_dir": str(stage),
    })
    return target


def run_browser_probe():
    """打印可用的渲染内核名，供 run_local.ps1 决定是否下载 Chromium。"""
    if not playwright_status():
        print("none")
        return 1
    try:
        from playwright.sync_api import sync_playwright
        with sync_playwright() as p:
            browser, label = launch_rendered_browser(p)
            browser.close()
        print(label)
        return 0
    except Exception:
        print("none")
        return 1


def run_selftest(args):
    """新环境自检：依赖、网络、翻译、可选官网可达性。返回退出码。"""
    import importlib

    line = "=" * 56
    print(line)
    print("enterprise-site-pipeline 环境自检")
    print(line)

    ok = True
    recoverable = []

    def probe(label, mod, required=True):
        nonlocal ok
        try:
            m = importlib.import_module(mod)
            ver = getattr(m, "__version__", "") or ""
            print(f"  [OK]   {label} {ver}".rstrip())
        except Exception as exc:
            tag = "FAIL" if required else "SKIP"
            print(f"  [{tag}] {label}：{exc}")
            if required:
                ok = False

    print("\n[1/4] Python 依赖")
    probe("openpyxl（Excel 读写）", "openpyxl")
    probe("python-docx（Word 生成）", "docx")
    probe("Pillow（视觉核对拼版）", "PIL")
    probe("playwright（JS 渲染，auto 模式使用）", "playwright", required=False)
    if playwright_status():
        try:
            from playwright.sync_api import sync_playwright
            with sync_playwright() as pw:
                browser, label = launch_rendered_browser(pw)
                browser.close()
            if label == "chromium":
                print("  [OK]   渲染内核：Playwright 内置 Chromium")
            else:
                print("  [OK]   渲染内核：本机 Microsoft Edge（无需下载 Chromium）")
        except Exception as exc:
            print(f"  [WARN] Playwright 已装但没有可用内核：{exc}")
            print("         本机装 Edge 即可自动启用；或运行： <python.exe> -m playwright install chromium")

    print("\n[2/4] 网络")
    for label, url in (("目标官网测试", "https://www.baidu.com/"),
                       ("MyMemory 翻译接口", TRANSLATE_URL + "?q=test&langpair=zh-CN|en")):
        r = fetch(url, args)
        if r.get("error"):
            if label.startswith("MyMemory"):
                print(f"  [WARN] {label}：{r['error']}（可恢复：Codex 用 --en 补英文草稿）")
                recoverable.append(f"{label}：{r['error']}")
            else:
                print(f"  [FAIL] {label}：{r['error']}")
                ok = False
        else:
            print(f"  [OK]   {label}（HTTP {r.get('status') or '?'}）")

    print("\n[3/4] 翻译连通性")
    got = translate_chunk("企业", args)
    if got:
        print(f"  [OK]   中译英可用：企业 → {got}")
    else:
        print("  [WARN] 中译英不可用（MyMemory 限流或网络不通）；"
              "不是环境致命错误，但默认门禁会阻止英文缺失时发布。")
        print("         必须由 Codex 基于中文事实生成 --en 兼容英文草稿，"
              "保留 定稿:false、自动翻译:true、备注:待人工核校；"
              "或由用户明确接受中文版。")
        recoverable.append("中译英不可用，需 Codex 补 --en 草稿")

    print("\n[4/4] 输入 Excel 与官网可达性")
    excel = Path(args.excel).expanduser() if args.excel else None
    if not excel or not excel.is_file():
        print("  [SKIP] 未提供 --excel，跳过输入与官网检查")
    else:
        try:
            companies = load_excel(excel)
        except Exception as exc:
            companies = []
            print(f"  [FAIL] 读取 Excel 失败：{exc}")
            ok = False
        print(f"  [OK]   读到 {len(companies)} 家企业")
        no_site = [c["name"] for c in companies if not c.get("website")]
        if no_site:
            print(f"  [INFO] {len(no_site)} 家没有官网列，将自动发现："
                  f"{'、'.join(no_site[:5])}{' 等' if len(no_site) > 5 else ''}")
        checked = 0
        for item in companies:
            if checked >= 3:
                break
            checked += 1
            if item.get("website"):
                page, r = fetch_site(norm_url(item["website"]), args)
                if page:
                    print(f"  [OK]   {item['name']}：官网可访问 {page.get('url')}")
                else:
                    print(f"  [FAIL] {item['name']}：官网不可访问 {r.get('error', '')}")
                    ok = False
            else:
                d = discover_site(item["name"], "", args)
                if d.get("url"):
                    print(f"  [OK]   {item['name']}：自动发现 {d['url']}"
                          f"（{d.get('confidence', '')}）")
                else:
                    print(f"  [WARN] {item['name']}：自动发现未确认官网"
                          f"（{d.get('error', '')[:120]}）")

    print("\n" + line)
    print("自检通过，可以运行流水线。" if ok else "自检未通过，请先修复上面标 FAIL 的项。")
    if recoverable:
        print("可恢复告警：" + "；".join(recoverable[:3]))
        print("处理方式：默认门禁不允许英文缺失直接发布；Codex 补 --en 草稿，或用户显式 --accept-no-english。")
    print(line)
    return 0 if ok else 1


def main():
    ap = argparse.ArgumentParser(description="本地企业官网资料包流水线")
    ap.add_argument("--excel", default="", help="包含企业名称的 Excel 文件")
    ap.add_argument("--selftest", action="store_true",
                    help="只做环境自检（依赖/网络/翻译/可选官网可达性）后退出，不跑流水线")
    ap.add_argument("--discover-only", dest="discover_only", action="store_true",
                    help="两阶段流程第一阶段：只做官网发现，产出「官网候选复核表.xlsx」后退出")
    ap.add_argument("--site-decisions", dest="site_decisions", default="",
                    help="官网复核表（--discover-only 产出）：决定优先；未复核的中/低置信度默认跳过")
    ap.add_argument("--out", default="enterprise-site-output", help="本地输出目录")
    ap.add_argument("--en", default="", help="可选：已确认的英文 JSON")
    ap.add_argument("--limit", type=int, default=0, help="只处理前 N 家，0=全部")
    ap.add_argument("--timeout", type=int, default=20, help="单次请求超时秒数")
    ap.add_argument("--max-bytes", type=int, default=8_000_000, help="单页最大下载字节")
    ap.add_argument("--max-image-bytes", type=int, default=5_000_000, help="单图最大下载字节")
    ap.add_argument("--max-pages", type=int, default=20, help="每家企业最多抓取页面数")
    ap.add_argument("--max-products", type=int, default=60, help="每家企业最多产品名数")
    ap.add_argument("--max-images", type=int, default=140, help="每家企业最多图片数")
    ap.add_argument("--search-candidates", type=int, default=8, help="官网发现候选域名数")
    ap.add_argument("--proxy", default="", help="可选 HTTP(S) 代理")
    ap.add_argument("--insecure", action="store_true", help="跳过 TLS 证书校验（仅测试环境）")
    ap.add_argument("--playwright", nargs="?", const="on", default="auto",
                    choices=["auto", "on", "off"],
                    help="Playwright 模式：auto=静态过薄时自动渲染（默认），on=强制渲染，off=只用静态")
    ap.add_argument("--html-dir", dest="html_dir", default="",
                    help="Codex 内置浏览器保存的离线 HTML 目录（含 manifest.json），网络/内核都不可用时兜底")
    ap.add_argument("--resources", default="",
                    help="用户资料总目录：每家企业一个任意命名子文件夹，或单家企业文件夹")
    ap.add_argument("--backup-dir", dest="backup_dir", default="",
                    help="原始资料备份目录（默认 <输出目录>\\原始资料备份，在 deliverable 之外）")
    ap.add_argument("--browser-probe", dest="browser_probe", action="store_true",
                    help="只探测渲染内核（chromium/msedge/msedge-exe/none）后退出")
    ap.add_argument("--require-visual", dest="require_visual", action="store_true", default=True,
                    help="兼容旧参数；视觉核对未完成默认按门禁 error 处理")
    ap.add_argument("--skip-visual-review", dest="skip_visual_review", action="store_true",
                    help="显式跳过视觉核对门禁（仅调试/用户明确授权时使用）")
    ap.add_argument("--publish-stage", default="",
                    help="不重新抓取，把已完成的 build/<run_id> 目录发布到 --out")
    ap.add_argument("--no-publish", dest="no_publish", action="store_true",
                    help="门禁通过也不发布，只保留 build/<run_id>（用于先做视觉核对）")
    ap.add_argument("--strict", action="store_true", help="门禁不通过时返回非零退出码")
    ap.add_argument("--no-translate", dest="no_translate", action="store_true",
                    help="关闭自动中英双语，仅输出中文并把英文层留空（仍需 --accept-no-english 才能发布）")
    ap.add_argument("--accept-no-english", dest="accept_no_english", action="store_true",
                    help="用户明确接受中文版；英文相关门禁降为告警，不建议用于默认双语交付")
    ap.add_argument("--translate-email", default="", help="可选：MyMemory 联系邮箱，用于提高匿名额度")
    ap.add_argument("--translate-delay", type=float, default=0.2, help="每次翻译调用后的间隔秒数")
    ap.add_argument("--allow-builder-cdn", dest="allow_builder_cdn", action="store_true",
                    help="放行建站平台自有 CDN（faiusr.com/faisys.com/508sys.com）的图片直链，"
                         "仅当来源页与官网同域时生效；默认关闭")
    ap.add_argument("--no-visual-review", dest="no_visual_review", action="store_true",
                    help="跳过图片拼版/核对表和视觉核对.json 生成")
    args = ap.parse_args()

    if args.browser_probe:
        return run_browser_probe()
    if args.publish_stage:
        stage = Path(args.publish_stage).expanduser().resolve()
        if not stage.is_dir():
            raise SystemExit(f"找不到已完成的构建目录：{stage}")
        manifest_path = stage / "manifest.json"
        manifest = json.loads(manifest_path.read_text(encoding="utf-8")) if manifest_path.is_file() else {}
        visual = stage / "视觉核对.json"
        passed, gate_log = run_gates(
            stage, stage / "deliverable", stage / "raw", stage / "en.json", stage / "汇总.xlsx",
            int(manifest.get("companies") or 0), visual if visual.is_file() else None,
            require_visual=not args.skip_visual_review,
            allow_no_english=bool(manifest.get("accept_no_english")),
            allow_builder_cdn=bool(args.allow_builder_cdn or manifest.get("allow_builder_cdn")),
        )
        log("\n" + (gate_log or "未运行门禁"))
        if not passed:
            log(f"门禁未通过，拒绝发布：{stage}")
            return 1
        out = Path(args.out).expanduser().resolve()
        published = publish(stage, out, stage.name)
        log(f"已发布：{published}")
        return 0
    if args.selftest:
        return run_selftest(args)
    args.html_overrides = load_html_overrides(args)
    if not args.excel:
        ap.error("缺少 --excel（仅 --selftest 模式可省略）")

    excel = Path(args.excel).expanduser().resolve()
    if not excel.is_file():
        raise SystemExit(f"找不到 Excel：{excel}")
    out = Path(args.out).expanduser().resolve()
    run_id = _dt.datetime.now().strftime("%Y%m%d-%H%M%S")
    companies = load_excel(excel)
    if args.limit:
        companies = companies[:args.limit]
    if not companies:
        raise SystemExit("Excel 中没有可处理的企业名称")
    if args.discover_only:
        return run_discover_only(companies, args, out)
    stage = out / "build" / run_id
    stage.mkdir(parents=True, exist_ok=True)
    log(f"输入 {excel}")
    log(f"企业 {len(companies)} 家 · 本地输出 {out} · run_id {run_id}")

    # 用户资料摄入：一个总目录，每家企业一个任意命名子文件夹。
    args.backup_dir = str(Path(args.backup_dir).expanduser().resolve()) if args.backup_dir \
        else str((out / "原始资料备份").resolve())
    resource_map = {}
    if args.resources:
        res_root = Path(args.resources).expanduser().resolve()
        if not res_root.is_dir():
            raise SystemExit(f"找不到资料目录：{res_root}")
        names = [c["name"] for c in companies]
        resource_map = match_resource_dir(names, res_root)
        # 单家企业：子目录匹配不到时，若总目录本身像企业文件夹则整体采用。
        if not resource_map and len(companies) == 1:
            resource_map[names[0]] = res_root
        matched = len(resource_map)
        log(f"资料目录 {res_root} · 匹配到 {matched}/{len(companies)} 家企业")
        if matched < len(companies):
            missing = [n for n in names if n not in resource_map]
            log(f"  未匹配到资料的企业（仅官网）：{', '.join(missing[:10])}")
    for item in companies:
        item["resource_dir"] = str(resource_map.get(item["name"], "")) if resource_map else ""

    decisions = {}
    if args.site_decisions:
        decisions_path = Path(args.site_decisions).expanduser()
        if not decisions_path.is_file():
            raise SystemExit(f"找不到官网复核表：{decisions_path}")
        decisions = load_site_decisions(decisions_path)
        log(f"官网复核表 {decisions_path} · 读到 {len(decisions)} 家企业决定")

    archives, skipped = [], []
    for i, item in enumerate(companies, 1):
        plan = plan_site(item, decisions, args)
        if plan["action"] == "skip":
            skipped.append({"名称": item["name"], "原因": plan.get("reason", "")})
            log(f"[{i}/{len(companies)}] {item['name']} —— 跳过：{plan.get('reason', '')}")
            continue
        item["_site_discovery"] = plan["discovery"]
        log(f"[{i}/{len(companies)}] {item['name']} ...")
        a = crawl_company(item, args, stage / "raw_home")
        if args.max_products:
            a["products"] = (a.get("products") or [])[:args.max_products]
        archives.append(a)
        log(f"  {a.get('status')} · 官网 {a.get('官网') or '未确认'} · 图片 {sum(len(a.get(k) or []) for k in IMAGE_DIRS)}")

    if skipped:
        write_json(stage / "官网复核结果.json", {
            "处理企业数": len(archives), "跳过企业数": len(skipped), "跳过": skipped,
        })
        log(f"跳过 {len(skipped)} 家：{'、'.join(x['名称'] for x in skipped[:10])}"
            f"{' 等' if len(skipped) > 10 else ''}"
            f"（未复核的中/低置信度或复核选择跳过）；详见 {stage / '官网复核结果.json'}")
    if not archives:
        log("所有企业均被跳过。请先运行 --discover-only 生成复核表，"
            "填写决定后再用 --site-decisions 重跑。")
        return 1 if args.strict else 0
    en_data = load_en(args.en, archives, args, out / "_translate_cache.json")
    backlog = english_backlog(archives, en_data)
    backlog_path = ""
    if backlog["企业"]:
        backlog_path = str(stage / "英文补译清单.json")
        write_json(Path(backlog_path), backlog)
        log(f"英文补译清单：{backlog_path}（默认门禁会阻止直接发布）")
    deliverable, raw_dir, en_path, summary = render_stage(stage, archives, en_data)
    visual_path, visual_log = None, ""
    if not args.no_visual_review:
        log("生成视觉核对材料（拼版 + 缩略图核对表 + 视觉核对.json）...")
        visual_path, visual_log = build_visual_review(stage, deliverable, raw_dir)
        if visual_path:
            log(f"  视觉核对材料：{stage / 'review'}")
            log("  请查看拼版后把结论写回 视觉核对.json，再跑 gates.py；未完成视觉核对会按 error 阻断门禁。")
        else:
            log(f"  视觉核对材料生成失败，门禁将按待核对处理：{visual_log.strip()[:300]}")
    write_json(stage / "manifest.json", {
        "run_id": run_id, "input": str(excel), "companies": len(archives),
        "excel": excel.name, "status_counts": dict(collections.Counter(a.get("status", "") for a in archives)),
        "english_final": bool(args.en),
        "english_mode": "final" if args.en else ("off" if args.no_translate else "auto"),
        "accept_no_english": bool(args.accept_no_english),
        "allow_builder_cdn": bool(args.allow_builder_cdn),
        "english_backlog": backlog_path,
        "visual_review": str(visual_path) if visual_path else "",
        "visual_review_dir": str(stage / "review") if visual_path else "",
        "generated_at": _dt.datetime.now().isoformat(timespec="seconds"),
    })
    passed, gate_log = run_gates(
        stage, deliverable, raw_dir, en_path, summary, len(archives), visual_path,
        require_visual=not args.skip_visual_review,
        allow_no_english=args.accept_no_english,
        allow_builder_cdn=args.allow_builder_cdn,
    )
    log("\n" + (gate_log or "未运行门禁"))
    if passed:
        if args.no_publish:
            log(f"门禁通过；--no-publish 已启用，未发布，构建目录：{stage}")
            return 0
        published = publish(stage, out, run_id)
        log(f"门禁通过，已发布：{published}")
        return 0
    log(f"产物已生成但门禁未全绿：{stage}")
    log("英文缺失或自动翻译失败时，必须补 --en <英文草稿/定稿.json>；"
        "只有用户明确接受中文版时才可加 --accept-no-english。")
    return 1 if args.strict else 0


if __name__ == "__main__":
    sys.exit(main())
